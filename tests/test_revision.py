"""End-to-end safety tests for editor-controlled manuscript revision."""

from __future__ import annotations

import io
import json
import uuid
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfgen import canvas
from sqlalchemy import create_engine, select
from sqlalchemy.dialects import postgresql
from sqlalchemy.schema import CreateTable
from sqlalchemy.orm import Session

from models import Base, Journal, Role, Submission, User, UserJournal
from services.core import AuthorizationError
from services.revision_analysis import analyze_manuscript, citation_tokens, classify_comment, extract_reviewer_comments, map_comment
from services.revision_docx import ApprovedPatch, UnsafeRevisionError, check_integrity, patch_manuscript
from services.revision_service import (
    RevisionRuleError, add_reviewer_file, analyze_revision_job, authorized_artifact_bytes,
    authorized_original_bytes, complete_revision_job, create_revision_job, decide_comment,
    draft_response_document, generate_revision_artifacts, revision_history,
)
from services.revision_ai import AIProposal, RevisionAIUnavailable
from services.revision_storage import LocalRevisionStorage, RevisionFileError, StorageUnavailable, SupabaseRevisionStorage, validate_revision_upload


def _save(document: Document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _manuscript() -> bytes:
    document = Document()
    document.sections[0].header.paragraphs[0].text = "JOURNAL HEADER"
    document.sections[0].footer.paragraphs[0].text = "JOURNAL FOOTER"
    document.add_heading("Controlled Revision of a Novel Method", 0)
    document.add_paragraph("A. Author, B. Author")
    document.add_paragraph("Example University")
    document.add_heading("Introduction", 1)
    target = document.add_paragraph()
    target.add_run("The novelty ").bold = True
    target.add_run("of this method is stated briefly. Kp = 2.55 [1]")
    document.add_paragraph("The discussion is concise and can be clearer.")
    document.add_heading("Methodology", 1)
    equation = document.add_paragraph("The model uses ")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x=2"
    run.append(text)
    math.append(run)
    equation._p.append(math)
    document.add_paragraph("Table I. Model parameters")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Ts"
    table.cell(1, 1).text = "0.03 s"
    document.add_paragraph("Figure 1. Experimental setup")
    image = bytes.fromhex("89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000b49444154789c636000020000050001a5f645400000000049454e44ae426082")
    document.add_picture(io.BytesIO(image))
    document.add_heading("References", 1)
    document.add_paragraph("[1] Verified existing source, 2024.")
    return _save(document)


def _pdf_review() -> bytes:
    output = io.BytesIO()
    pdf = canvas.Canvas(output)
    pdf.drawString(60, 750, "1. Clarify the novelty in the Introduction.")
    pdf.save()
    return output.getvalue()


def test_extraction_classification_and_mapping():
    structure = analyze_manuscript(_manuscript())
    assert structure.title.startswith("Controlled Revision")
    assert structure.table_count == structure.figure_count == structure.equation_count == 1
    assert len(structure.references) == 1
    first = extract_reviewer_comments("r1.txt", b"Major Comments\n1. Clarify the novelty in the Introduction.\n2. Please add a new experiment.", "Reviewer 1")
    assert len(first) == 2
    assert first[0].raw_text == "Clarify the novelty in the Introduction."
    assert first[0].severity == "MAJOR"
    assert first[1].category == "DATA_CHANGE" and first[1].author_input_required
    assert map_comment(first[0], structure).paragraph_id
    assert len(extract_reviewer_comments("r2.pdf", _pdf_review(), "Reviewer 2")) == 1
    reviewer_doc = Document()
    reviewer_doc.add_paragraph("1. Clarify the novelty in the Introduction.")
    assert len(extract_reviewer_comments("r3.docx", _save(reviewer_doc), "Reviewer 3")) == 1
    assert classify_comment("Please add a DOI reference")[3]
    assert classify_comment("Please replace Figure 4")[3]
    assert citation_tokens("Smith et al. (2020) and (Jones, 2021), doi:10.1234/example") != citation_tokens(
        "Smith et al. (2020) and (Jones, 2022), doi:10.1234/example")


def test_targeted_docx_patch_preserves_package_and_science():
    original = _manuscript()
    structure = analyze_manuscript(original)
    target = next(item for item in structure.paragraphs if item.text.startswith("The novelty"))
    replacement = "The novelty of this method is clearly distinguished from prior work. Kp = 2.55 [1]"
    patch = ApprovedPatch(target.identifier, target.text, replacement)
    clean = patch_manuscript(original, [patch], highlight=False)
    marked = patch_manuscript(original, [patch], highlight=True)
    assert check_integrity(original, clean, [patch]).passed
    assert check_integrity(original, marked, [patch]).passed
    assert target.text not in "\n".join(item.text for item in analyze_manuscript(clean).paragraphs)
    with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(marked)) as after:
        assert before.read("word/header1.xml") == after.read("word/header1.xml")
        assert before.read("word/footer1.xml") == after.read("word/footer1.xml")
        assert before.read("word/media/image1.png") == after.read("word/media/image1.png")
        xml = after.read("word/document.xml")
        assert b'm:atr' not in xml
        assert b"<m:oMath>" in xml or b"m:oMath" in xml
        assert b'w:val="yellow"' in xml
    green = patch_manuscript(original, [patch], highlight=True, highlight_color="green")
    with ZipFile(io.BytesIO(green)) as archive:
        assert b'w:val="green"' in archive.read("word/document.xml")
    paragraph = next(p for p in Document(io.BytesIO(marked)).paragraphs if p.text == replacement)
    assert paragraph.runs[0].bold is True
    with pytest.raises(UnsafeRevisionError):
        patch_manuscript(original, [ApprovedPatch(target.identifier, target.text, replacement.replace("2.55", "1.6"))], highlight=False)
    with pytest.raises(UnsafeRevisionError):
        patch_manuscript(original, [ApprovedPatch(target.identifier, target.text, replacement.replace("[1]", "[2]"))], highlight=False)
    authorized_number_patch = ApprovedPatch(target.identifier, target.text, replacement.replace("2.55", "1.6"),
                                             "Author confirmed corrected parameter")
    approved_numeric = patch_manuscript(original, [authorized_number_patch], highlight=False)
    numeric_report = check_integrity(original, approved_numeric, [authorized_number_patch])
    assert numeric_report.passed and len(numeric_report.checks["approved_numeric_changes"]) == 1


def test_exact_text_offsets_numbered_references_and_page_break(tmp_path: Path):
    document = Document(io.BytesIO(_manuscript()))
    references_heading = next(paragraph for paragraph in document.paragraphs if paragraph.text == "References")
    references_heading.text = "5 References"
    leading = document.add_paragraph(" Leading and trailing ")
    references_heading._p.addprevious(leading._p)
    page_break = OxmlElement("w:p")
    page_break_run = OxmlElement("w:r")
    break_node = OxmlElement("w:br")
    break_node.set(qn("w:type"), "page")
    page_break_run.append(break_node)
    page_break.append(page_break_run)
    references_heading._p.addprevious(page_break)
    original = _save(document)
    (tmp_path / "multipage_original.docx").write_bytes(original)
    structure = analyze_manuscript(original)
    assert structure.references == ("[1] Verified existing source, 2024.",)
    target = next(p for p in structure.paragraphs if "Leading and trailing" in p.text)
    assert target.text == " Leading and trailing "
    patch = ApprovedPatch(target.identifier, target.text, " Leading and corrected trailing ")
    revised = patch_manuscript(original, [patch], highlight=True)
    (tmp_path / "multipage_revised.docx").write_bytes(revised)
    assert check_integrity(original, revised, [patch]).passed
    assert next(p for p in analyze_manuscript(revised).paragraphs if p.identifier == target.identifier).text == patch.proposed_text


@pytest.fixture()
def case(tmp_path: Path):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    storage = LocalRevisionStorage(tmp_path / "private")
    with Session(engine, expire_on_commit=False) as session:
        journal = Journal(name="Journal A", abbreviation="JA")
        other = Journal(name="Journal B", abbreviation="JB")
        editor = User(email="editor@local.test", display_name="Editor", password_hash="x", role=Role.JOURNAL_ADMIN)
        outsider = User(email="outsider@local.test", display_name="Outsider", password_hash="x", role=Role.JOURNAL_ADMIN)
        viewer = User(email="viewer@local.test", display_name="Viewer", password_hash="x", role=Role.FINANCE)
        session.add_all([journal, other, editor, outsider, viewer])
        session.flush()
        session.add_all([UserJournal(user_id=editor.id, journal_id=journal.id), UserJournal(user_id=outsider.id, journal_id=other.id)])
        submission = Submission(journal_id=journal.id, ojs_submission_id="11397", manuscript_title="Controlled Revision of a Novel Method",
                                corresponding_author="A. Author", email="a@local.test", created_by=editor.id)
        session.add(submission)
        session.commit()
        yield session, storage, journal, other, editor, outsider, viewer, submission
    engine.dispose()


def test_revision_end_to_end_two_reviewers_rounds_auth_and_failure(case, monkeypatch):
    import services.revision_service as revision_service

    monkeypatch.setattr(revision_service, "settings", SimpleNamespace(
        revision_allow_full_polishing=False, revision_ai_enabled=True, revision_ai_provider="test",
        revision_highlight_style="yellow"))

    class SafeAI:
        def generate_revision_proposal(self, **kwargs):
            return AIProposal(kwargs["paragraph_id"], kwargs["original_text"],
                kwargs["original_text"].replace("stated briefly", "explained in relation to prior work"),
                "Clarifies the novelty without changing data", 95, False)

    session, storage, journal, other, editor, outsider, viewer, submission = case
    original = _manuscript()
    job = create_revision_job(session, journal, editor, submission=submission, article_title=submission.manuscript_title,
        submission_identifier="11397", revision_round=1, mode="REVIEWER_DRIVEN", manuscript_filename="original.docx",
        manuscript_content=original, storage=storage)
    assert authorized_original_bytes(session, journal, editor, job.id, storage=storage)[1] == original
    add_reviewer_file(session, journal, editor, job.id, reviewer_label="Reviewer 1", filename="r1.txt",
                      content=b"1. Clarify the novelty in the Introduction.", storage=storage)
    add_reviewer_file(session, journal, editor, job.id, reviewer_label="Reviewer 2", filename="r2.pdf",
                      content=_pdf_review(), storage=storage)
    comments = analyze_revision_job(session, journal, editor, job.id, ai_consent=True,
                                    ai_service=SafeAI(), storage=storage)
    assert len(comments) == 2 and {c.reviewer_label for c in comments} == {"Reviewer 1", "Reviewer 2"}
    assert all(c.status == "PROPOSED" and c.proposed_text for c in comments)
    draft_text = "\n".join(paragraph.text for paragraph in Document(io.BytesIO(draft_response_document(session, journal, editor, job.id))).paragraphs)
    assert "Not Final" in draft_text and "Pending author input" in draft_text
    target = next(p for p in analyze_manuscript(original).paragraphs if p.text.startswith("The novelty"))
    with pytest.raises(RevisionRuleError):
        decide_comment(session, journal, editor, job.id, comments[0].id, action="ACCEPTED", paragraph_id=target.identifier,
                       proposed_text=target.text.replace("2.55", "1.6"), storage=storage)
    replacement = target.text.replace("stated briefly", "explained in relation to prior work")
    decide_comment(session, journal, editor, job.id, comments[0].id, action="EDITED", paragraph_id=target.identifier,
                   proposed_text=replacement, reason="Clarifies novelty", storage=storage)
    with pytest.raises(RevisionRuleError, match="already targets"):
        decide_comment(session, journal, editor, job.id, comments[1].id, action="ACCEPTED", paragraph_id=target.identifier,
                       proposed_text=replacement, storage=storage)
    decide_comment(session, journal, editor, job.id, comments[1].id, action="REJECTED", reason="Duplicate of Reviewer 1's point", storage=storage)
    artifacts = generate_revision_artifacts(session, journal, editor, job.id, storage=storage)
    assert {a.kind for a in artifacts} == {"REVISED_MANUSCRIPT", "CLEAN_MANUSCRIPT", "RESPONSE_TO_REVIEWERS", "REVISION_LOG"}
    for artifact in artifacts:
        filename, payload = authorized_artifact_bytes(session, journal, editor, artifact.id, storage=storage)
        assert filename == artifact.filename and payload
    assert authorized_original_bytes(session, journal, editor, job.id, storage=storage)[1] == original
    assert any(event.action == "REVISION_PROPOSAL_EDITED" for event in revision_history(session, journal, editor, job.id))
    with pytest.raises(AuthorizationError):
        authorized_artifact_bytes(session, journal, outsider, artifacts[0].id, storage=storage)
    with pytest.raises(AuthorizationError):
        authorized_original_bytes(session, journal, viewer, job.id, storage=storage)
    complete_revision_job(session, journal, editor, job.id)
    second = create_revision_job(session, journal, editor, submission=submission, article_title=submission.manuscript_title,
        submission_identifier="11397", revision_round=2, mode="CONSERVATIVE", manuscript_filename="round2.docx",
        manuscript_content=original, storage=storage)
    assert second.id != job.id and second.revision_round == 2
    with pytest.raises(RevisionRuleError):
        create_revision_job(session, journal, editor, submission=submission, article_title=submission.manuscript_title,
            submission_identifier="11397", revision_round=2, mode="CONSERVATIVE", manuscript_filename="duplicate.docx",
            manuscript_content=original, storage=storage)


def test_upload_guards_and_postgres_ddl():
    with pytest.raises(RevisionFileError, match="PDF manuscript"):
        validate_revision_upload("paper.pdf", _pdf_review(), manuscript=True)
    with pytest.raises(RevisionFileError):
        validate_revision_upload("bad.docx", b"not a zip", manuscript=True)
    for table in Base.metadata.sorted_tables:
        if table.name.startswith("revision_"):
            ddl = str(CreateTable(table).compile(dialect=postgresql.dialect()))
            assert table.name in ddl


def test_private_supabase_storage_adapter(monkeypatch):
    requests = []
    objects = {}

    class Response:
        def __init__(self, payload):
            self.payload = payload

        def __enter__(self):
            return self

        def __exit__(self, *args):
            return False

        def read(self, limit):
            return self.payload[:limit]

    public = False

    def fake_open(request, timeout):
        requests.append(request)
        if "/bucket/" in request.full_url:
            return Response(json.dumps({"public": public}).encode())
        if request.get_method() == "POST":
            objects[request.full_url.split("/object/")[1]] = request.data
            return Response(b'{"Key":"ok"}')
        key = request.full_url.split("/object/authenticated/")[1]
        return Response(objects[key])

    monkeypatch.setattr("urllib.request.urlopen", fake_open)
    storage = SupabaseRevisionStorage("https://example.supabase.co", "secret-server-key", "jas-private-revisions")
    key, digest = storage.put(uuid.uuid4(), "original", b"private manuscript", ".docx")
    assert storage.read(key) == b"private manuscript"
    assert any("/object/authenticated/" in req.full_url for req in requests)
    assert not any("/object/public/" in req.full_url for req in requests)
    public = True
    with pytest.raises(StorageUnavailable, match="PRIVATE"):
        storage.put(uuid.uuid4(), "original", b"blocked", ".docx")


def test_ai_failure_is_graceful_and_full_polish_is_editorial(case, monkeypatch):
    import services.revision_service as revision_service

    session, storage, journal, _, editor, _, _, submission = case
    monkeypatch.setattr(revision_service, "settings", SimpleNamespace(
        revision_allow_full_polishing=True, revision_ai_enabled=True, revision_ai_provider="test",
        revision_highlight_style="yellow"))
    job = create_revision_job(session, journal, editor, submission=submission, article_title=submission.manuscript_title,
        submission_identifier="11397", revision_round=1, mode="FULL_ACADEMIC_POLISHING",
        manuscript_filename="original.docx", manuscript_content=_manuscript(), storage=storage)
    add_reviewer_file(session, journal, editor, job.id, reviewer_label="Reviewer 1", filename="r1.txt",
                      content=b"1. Clarify the novelty in the Introduction.", storage=storage)

    class FailingAI:
        def generate_revision_proposal(self, **kwargs):
            raise RevisionAIUnavailable("offline")

    comments = analyze_revision_job(session, journal, editor, job.id, ai_consent=True,
                                    ai_service=FailingAI(), storage=storage)
    assert len(comments) >= 2
    assert any(c.source_type == "EDITORIAL_POLISHING" for c in comments)
    assert all(c.status == "EDITOR_CONFIRMATION_REQUIRED" for c in comments)
    assert len(revision_service._response_rows(job)) == 1


def test_ai_scientific_number_change_is_discarded(case, monkeypatch):
    import services.revision_service as revision_service

    session, storage, journal, _, editor, _, _, submission = case
    monkeypatch.setattr(revision_service, "settings", SimpleNamespace(
        revision_allow_full_polishing=False, revision_ai_enabled=True, revision_ai_provider="test",
        revision_highlight_style="yellow"))

    class UnsafeAI:
        def generate_revision_proposal(self, **kwargs):
            return AIProposal(kwargs["paragraph_id"], kwargs["original_text"],
                kwargs["original_text"].replace("2.55", "1.6"), "Change result", 99, False)

    job = create_revision_job(session, journal, editor, submission=submission, article_title=submission.manuscript_title,
        submission_identifier="11397", revision_round=1, mode="REVIEWER_DRIVEN",
        manuscript_filename="original.docx", manuscript_content=_manuscript(), storage=storage)
    add_reviewer_file(session, journal, editor, job.id, reviewer_label="Reviewer 1", filename="r1.txt",
                      content=b"1. Clarify the novelty in the Introduction.", storage=storage)
    comments = analyze_revision_job(session, journal, editor, job.id, ai_consent=True,
                                    ai_service=UnsafeAI(), storage=storage)
    assert comments[0].status == "AUTHOR_INPUT_REQUIRED"
    assert comments[0].proposed_text is None
