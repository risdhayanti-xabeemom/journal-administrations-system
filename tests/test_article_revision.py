"""Template-only formatting, data isolation and document-integrity regression."""

from __future__ import annotations

import copy
import io
import json
import uuid
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.dialects import postgresql
from sqlalchemy.orm import Session
from sqlalchemy.schema import CreateTable

from models import ArticleRevisionArtifact, ArticleRevisionJob, Base, Journal, Role, Submission, User, UserJournal
from services.article_formatting import (
    ArticleFormattingError, apply_safe_fixes, audit_article, compliance_report_xlsx,
    extract_style_profile, formatting_integrity,
)
from services.article_metadata import (
    ArticleMetadataError, _xml_placeholders, elkolind_master_warnings,
    graft_article_master_header_footer, graft_integrity, resolve_elkolind_metadata, short_title_4w,
)
from services.article_revision_service import (
    ArticleRevisionRuleError, accept_all_low_risk_review_findings, approve_article_fixes,
    article_revision_history, article_revision_jobs,
    audit_article_revision_job, auto_format_and_generate,
    authorized_article_artifact_bytes, complete_article_revision_job, create_article_revision_job,
    generate_formatted_manuscript, job_auto_format_result, job_fast_findings, job_review_decisions,
    job_resolution_summary, resolve_article_review_findings,
    select_active_article_template_for_job, update_article_job_metadata,
)
from services.article_template_service import (
    ELKOLIND_INITIAL_RULES,
    activate_article_template, active_article_template, article_template_config,
    article_template_versions, authorized_article_template_bytes, update_article_rules,
    upload_article_template,
)
from services.core import AuthorizationError
from services.revision_analysis import analyze_manuscript
from services.revision_storage import LocalRevisionStorage
from services.revision_storage import validate_revision_upload, RevisionFileError
from test_revision import _manuscript, _save


def _template(font: str = "Gadugi") -> bytes:
    document = Document(io.BytesIO(_manuscript()))
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.8)
    document.styles["Normal"].font.name = font
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Title"].font.name = font
    document.styles["Title"].font.size = Pt(18)
    document.styles["Heading 1"].font.name = font
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Caption"].font.name = font
    document.styles["Caption"].font.size = Pt(9)
    return _save(document)


def _elkolind_template() -> bytes:
    document = Document(io.BytesIO(_template()))
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.different_first_page_header_footer = True
    for header in (section.header, section.even_page_header, section.first_page_header):
        header.paragraphs[0].text = (
            "Jurnal Elkolind Volume {{volume}}, Nomor {{issue}}, "
            "{{publication_month}} {{publication_year}}"
        )
        header.add_paragraph("DOI: {{doi_full}}")
    with ZipFile(io.BytesIO(_manuscript())) as source:
        image_name = next(name for name in source.namelist() if name.startswith("word/media/"))
        barcode_image = source.read(image_name)
    section.footer.paragraphs[0].text = "p-ISSN: 2356-0533; e-ISSN: 2355-9195"
    section.footer.paragraphs[0].add_run().add_picture(io.BytesIO(barcode_image), width=Inches(0.2))
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    section.footer.paragraphs[0]._p.append(page)
    section.first_page_footer.paragraphs[0].text = "p-ISSN: 2356-0533; e-ISSN: 2355-9195"
    section.even_page_footer.paragraphs[0].text = (
        "{{first_author}}: {{short_title_4w}}…    "
        "p-ISSN: 2356-0533; e-ISSN: 2355-9195"
    )
    return _save(document)


def _noncompliant_elkolind_manuscript() -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(12)
    document.styles["Heading 1"].font.name = "Cambria"
    document.styles["Heading 1"].font.size = Pt(24)
    document.styles["Heading 2"].font.name = "Aptos Display"
    document.styles["Heading 2"].font.size = Pt(16)

    title = document.add_heading("Analisis Kendali Motor dengan Metode Baru", 0)
    title.runs[0].font.name = "Times New Roman"
    title.runs[0].font.size = Pt(16)
    authors = document.add_paragraph()
    authors.add_run("A. Author, ").font.name = "Times New Roman"
    authors.add_run("B. Author").font.name = "Cambria"
    document.add_paragraph("Example University").runs[0].font.name = "Aptos"
    corresponding = document.add_paragraph("* Corresponding author: author@example.test")
    corresponding.runs[0].italic = True
    corresponding.runs[0].font.name = "Times New Roman"

    abstract_heading = document.add_heading("ABSTRAK", 1)
    abstract_heading.runs[0].font.name = "Cambria"
    abstract_heading.runs[0].font.size = Pt(24)
    abstract = document.add_paragraph()
    abstract_start = abstract.add_run("Pengujian menggunakan ")
    abstract_start.font.name = "Cambria"
    abstract_start.bold = True
    abstract_result = abstract.add_run("Kp = 2.55 dan menghasilkan akurasi 95%.")
    abstract_result.font.name = "Times New Roman"
    abstract_result.bold = True
    keywords = document.add_paragraph("Kata Kunci: kendali; motor; pengujian")
    keywords.runs[0].font.name = "Cambria"
    keywords.runs[0].bold = True

    heading1 = document.add_heading("1. PENDAHULUAN", 1)
    heading1.runs[0].font.name = "Cambria"
    heading1.runs[0].font.size = Pt(24)
    body = document.add_paragraph()
    body.add_run("Teks utama menggunakan ").font.name = "Times New Roman"
    emphasized = body.add_run("campuran format")
    emphasized.font.name = "Aptos"
    emphasized.bold = True
    body.add_run(" dengan Ts = 0.03 s [1].").font.name = "Cambria"

    narrative_figure = document.add_paragraph(
        "Rangkaian pada Gambar 1 menggunakan ESP32 dan sensor ultrasonik."
    )
    narrative_figure.runs[0].font.name = "Times New Roman"
    narrative_figure.runs[0].bold = True
    narrative_table = document.add_paragraph("Hasil pengukuran ditunjukkan pada Tabel I.")
    narrative_table.runs[0].font.name = "Cambria"
    narrative_table.runs[0].bold = True

    heading2 = document.add_heading("2.1 Format Halaman", 2)
    heading2.runs[0].font.name = "Aptos Display"
    heading2.runs[0].font.size = Pt(16)
    equation = document.add_paragraph("Persamaan dipertahankan ")
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = "x=2"
    math_run.append(math_text)
    math.append(math_run)
    equation._p.append(math)

    image = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
        "0000000b49444154789c636000020000050001a5f645400000000049454e44ae426082"
    )
    document.add_picture(io.BytesIO(image))
    document.add_paragraph("Gambar 1 : Grafik hasil pengujian")
    document.add_picture(io.BytesIO(image))
    document.add_paragraph("Figure 2: Konfigurasi sistem")

    document.add_paragraph("Tabel 1: hasil pengukuran tegangan")
    table1 = document.add_table(rows=2, cols=2)
    table1.cell(0, 0).text = "Parameter"
    table1.cell(0, 1).text = "Nilai"
    table1.cell(1, 0).text = "Tegangan"
    table1.cell(1, 1).text = "220 V"
    document.add_paragraph("TABEL II: perbandingan metode")
    table2 = document.add_table(rows=2, cols=2)
    table2.cell(0, 0).text = "Metode"
    table2.cell(0, 1).text = "Akurasi"
    table2.cell(1, 0).text = "Usulan"
    table2.cell(1, 1).text = "95%"

    document.add_heading("3. HASIL DAN PEMBAHASAN", 1)
    document.add_paragraph("Nilai ilmiah tetap Kp = 2.55, Ts = 0.03 s, dan 95%.")
    document.add_heading("DAFTAR PUSTAKA", 1)
    reference = document.add_paragraph("[1] Verified source, 2024. doi:10.1234/example")
    reference.runs[0].font.name = "Times New Roman"
    reference.runs[0].font.size = Pt(11)
    return _save(document)


def _manuscript_with_abstract_keywords() -> bytes:
    """Representative manuscript with required metadata sections for generation-gate tests."""
    document = Document(io.BytesIO(_manuscript()))
    introduction = next(paragraph for paragraph in document.paragraphs if paragraph.text == "Introduction")
    table_caption = next(paragraph for paragraph in document.paragraphs
                         if paragraph.text == "Table I. Model parameters")
    table_caption.text = "Tabel I: hasil pengujian"
    abstract_heading = document.add_paragraph("ABSTRAK", style="Heading 1")
    abstract = document.add_paragraph(" ".join(
        ["Penelitian ini mempertahankan data ilmiah dan memeriksa kepatuhan format secara deterministik."] * 12
    ))
    keywords = document.add_paragraph("Kata Kunci: kendali; format; integritas")
    introduction._p.addprevious(abstract_heading._p)
    introduction._p.addprevious(abstract._p)
    introduction._p.addprevious(keywords._p)
    return _save(document)


def _append_protected_caption(document: Document, before_table_text: str, field_text: str,
                              after_table_text: str, bookmark_id: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(before_table_text)
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), bookmark_id)
    bookmark_start.set(qn("w:name"), f"caption_{bookmark_id}")
    paragraph._p.append(bookmark_start)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "SEQ Table")
    field_run = OxmlElement("w:r")
    field_value = OxmlElement("w:t")
    field_value.text = field_text
    field_run.append(field_value)
    field.append(field_run)
    paragraph._p.append(field)
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), bookmark_id)
    paragraph._p.append(bookmark_end)
    paragraph.add_run(after_table_text)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metode"
    table.cell(0, 1).text = "Galat"
    table.cell(1, 0).text = "Usulan"
    table.cell(1, 1).text = "2.55%"


def _protected_caption_manuscript() -> bytes:
    document = Document()
    document.add_heading("Analisis Kendali Motor", 0)
    document.add_paragraph("A. Author")
    document.add_paragraph("Example University")
    document.add_heading("ABSTRAK", 1)
    document.add_paragraph(" ".join(
        ["Penelitian ini mempertahankan data ilmiah dan memeriksa kepatuhan format secara deterministik."] * 12
    ))
    document.add_paragraph("Kata Kunci: kendali; format; integritas")
    document.add_heading("1. PENDAHULUAN", 1)
    document.add_paragraph("Nilai Kp = 2.55 dan akurasi 95% dipertahankan.")
    _append_protected_caption(document, "TABEL ", "II", " : PERBANDINGAN METODE", "41")
    _append_protected_caption(document, "Tabel ", "2", ": Perbandingan metode", "42")
    document.add_heading("DAFTAR PUSTAKA", 1)
    document.add_paragraph("[1] Verified source, 2024. doi:10.1234/example")
    return _save(document)


def _paragraph_xml_by_text(content: bytes, text: str) -> bytes:
    with ZipFile(io.BytesIO(content)) as package:
        root = etree.fromstring(package.read("word/document.xml"))
    paragraph = next(node for node in root.xpath(".//w:body//w:p", namespaces={"w": qn("w:p")[1:].split("}")[0]})
                     if "".join(node.xpath(".//w:t/text()", namespaces={"w": qn("w:p")[1:].split("}")[0]})) == text)
    return etree.tostring(paragraph)


def _split_placeholder_runs(root: etree._Element, tokens: tuple[str, ...]) -> None:
    """Split placeholders across cloned Word runs without changing visible text."""
    word = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    for token in tokens:
        for node in list(root.xpath(".//w:t", namespaces={"w": word[1:-1]})):
            text = node.text or ""
            position = text.find(token)
            if position < 0:
                continue
            run = node.getparent()
            parent = run.getparent() if run is not None else None
            assert run is not None and run.tag == word + "r" and parent is not None
            assert len(run.xpath("./w:t", namespaces={"w": word[1:-1]})) == 1
            midpoint = len(token) // 2
            node.text = text[:position] + token[:midpoint]
            cloned_run = copy.deepcopy(run)
            cloned_text = cloned_run.find(word + "t")
            cloned_text.text = token[midpoint:] + text[position + len(token):]
            parent.insert(parent.index(run) + 1, cloned_run)
            for text_node in (node, cloned_text):
                if (text_node.text or "").startswith(" ") or (text_node.text or "").endswith(" "):
                    text_node.set(xml_space, "preserve")
                else:
                    text_node.attrib.pop(xml_space, None)


def _elkolind_textbox_template(*, split_runs: bool = False) -> bytes:
    """Valid Word regression master whose placeholders live in textboxes."""
    source = (Path(__file__).parent / "fixtures" / "elkolind_textbox_master.docx").read_bytes()
    if not split_runs:
        return source
    with ZipFile(io.BytesIO(source)) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
        members = archive.infolist()
    tokens = ("{{volume}}", "{{issue}}", "{{publication_month}}", "{{publication_year}}",
              "{{doi_suffix}}", "{{first_author}}", "{{short_title_4w}}")
    for name in sorted(name for name in files if name.startswith("word/") and
                       any(kind in Path(name).name.lower() for kind in ("header", "footer")) and
                       name.endswith(".xml")):
        root = etree.fromstring(files[name])
        _split_placeholder_runs(root, tokens)
        files[name] = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    output = io.BytesIO()
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        for member in members:
            archive.writestr(member, files[member.filename])
    return output.getvalue()


@pytest.fixture()
def case(tmp_path: Path):
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    storage = LocalRevisionStorage(tmp_path / "private")
    with Session(engine, expire_on_commit=False) as session:
        ja = Journal(name="ELKOLIND test journal", abbreviation="ELKOLIND")
        jb = Journal(name="JASENS test journal", abbreviation="JASENS")
        editor = User(email="article-editor@local.test", display_name="Editor", password_hash="x", role=Role.JOURNAL_ADMIN)
        outsider = User(email="article-outsider@local.test", display_name="Other", password_hash="x", role=Role.JOURNAL_ADMIN)
        finance = User(email="article-finance@local.test", display_name="Finance", password_hash="x", role=Role.FINANCE)
        session.add_all([ja, jb, editor, outsider, finance])
        session.flush()
        session.add_all([UserJournal(user_id=editor.id, journal_id=ja.id),
                         UserJournal(user_id=outsider.id, journal_id=jb.id)])
        submission = Submission(journal_id=ja.id, ojs_submission_id="11397",
            manuscript_title="Controlled Revision of a Novel Method", corresponding_author="A. Author",
            email="author@local.test", created_by=editor.id)
        session.add(submission)
        session.commit()
        yield session, storage, ja, jb, editor, outsider, finance, submission
    engine.dispose()


def test_profile_audit_and_safe_patching_preserve_science():
    original = _manuscript()
    profile = extract_style_profile(_template())
    rules = {"required_sections": ["Introduction", "Methodology", "Conclusion", "References"],
             "title_max_words": 30, "figure_caption_prefix": "Figure", "table_caption_prefix": "Table"}
    findings = audit_article(original, profile, rules)
    assert any(item.check == "Page Width" and item.status == "SAFE_FIX_AVAILABLE" for item in findings)
    assert any(item.category == "Body" and item.attribute == "font" and item.status == "SAFE_FIX_AVAILABLE" for item in findings)
    assert any(item.status == "MISSING_REQUIRED_SECTION" and "Conclusion" in item.check for item in findings)
    safe = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    formatted = apply_safe_fixes(original, findings, safe)
    report = formatting_integrity(original, formatted)
    assert report["passed"], report
    assert all(report["checks"].values())
    assert Document(io.BytesIO(formatted)).sections[0].page_width.twips == profile["page"]["page_width"]
    with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(formatted)) as after:
        assert set(before.namelist()) == set(after.namelist())
        for name in before.namelist():
            if name != "word/document.xml":
                assert before.read(name) == after.read(name)
        assert b"2.55" in after.read("word/document.xml")
        assert b"0.03 s" in after.read("word/document.xml")
    with pytest.raises(ArticleFormattingError):
        apply_safe_fixes(original, findings, {next(item.id for item in findings if item.status == "MISSING_REQUIRED_SECTION")})


def test_elkolind_deterministic_typography_caption_and_run_normalization():
    original = _noncompliant_elkolind_manuscript()
    profile = extract_style_profile(_elkolind_template())
    rules = {**ELKOLIND_INITIAL_RULES, "left_margin_mm": 14.32, "right_margin_mm": 14.32}
    findings = audit_article(original, profile, rules)
    semantic = next(item for item in findings
                    if item.check == "ELKOLIND deterministic typography and caption formatting")
    assert semantic.status == "SAFE_FIX_AVAILABLE"
    table_caption_fix = next(item for item in findings
                             if item.check == "Caption notation" and
                             item.detected == "Tabel 1: hasil pengukuran tegangan")
    assert table_caption_fix.status == "SAFE_FIX_AVAILABLE"
    selected = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    formatted = apply_safe_fixes(original, findings, selected)
    integrity = formatting_integrity(original, formatted)
    assert integrity["passed"], integrity
    assert integrity["checks"]["numbers_unchanged"]
    assert integrity["checks"]["equations_unchanged"]
    assert integrity["checks"]["tables_unchanged"]

    post = audit_article(formatted, profile, rules)
    assert next(item for item in post
                if item.check == "ELKOLIND deterministic typography and caption formatting").status == "COMPLIANT"
    expected_summaries = {"Global font", "Abstract font", "Abstract body bold", "Body font",
                          "Body paragraph-wide bold", "Heading 1", "Heading 2",
                          "Figure caption format", "Table caption format"}
    assert {item.check for item in post if item.category == "ELKOLIND" and item.status == "COMPLIANT"} == expected_summaries
    assert all(item.status == "COMPLIANT" for item in post
               if item.check == "Caption notation")

    document = Document(io.BytesIO(formatted))
    paragraphs = {paragraph.text: paragraph for paragraph in document.paragraphs}
    assert "Gambar 1: Grafik hasil pengujian" in paragraphs
    assert "Gambar 2: Konfigurasi sistem" in paragraphs
    assert "TABEL I : HASIL PENGUKURAN TEGANGAN" in paragraphs
    assert "TABEL II : PERBANDINGAN METODE" in paragraphs
    assert paragraphs["3. HASIL DAN PEMBAHASAN"].runs[0].font.size.pt == 10
    assert paragraphs["3. HASIL DAN PEMBAHASAN"].runs[0].bold is True
    assert paragraphs["ABSTRAK"].runs[0].font.size.pt == 9
    assert paragraphs["ABSTRAK"].runs[0].bold is True
    assert paragraphs["* Corresponding author: author@example.test"].runs[0].italic is True
    assert "Kp = 2.55" in "\n".join(paragraphs)
    assert "Ts = 0.03 s" in "\n".join(paragraphs)
    assert "doi:10.1234/example" in "\n".join(paragraphs)
    assert all(run.font.name == "Gadugi" and run.font.size.pt == 9
               for run in paragraphs["Pengujian menggunakan Kp = 2.55 dan menghasilkan akurasi 95%."].runs)
    assert all(run.bold is False
               for run in paragraphs["Pengujian menggunakan Kp = 2.55 dan menghasilkan akurasi 95%."].runs)
    assert paragraphs["Pengujian menggunakan Kp = 2.55 dan menghasilkan akurasi 95%."].alignment == 3
    assert not all(run.bold is True for run in paragraphs["Kata Kunci: kendali; motor; pengujian"].runs)
    assert all(run.font.name == "Gadugi" and run.font.size.pt == 10
               for run in paragraphs["Teks utama menggunakan campuran format dengan Ts = 0.03 s [1]."].runs)
    for narrative in (
        "Rangkaian pada Gambar 1 menggunakan ESP32 dan sensor ultrasonik.",
        "Hasil pengukuran ditunjukkan pada Tabel I.",
    ):
        assert narrative in paragraphs
        assert all(run.font.name == "Gadugi" and run.font.size.pt == 10 and run.bold is False
                   for run in paragraphs[narrative].runs)
        assert paragraphs[narrative].alignment == 3
    assert all(run.font.name == "Gadugi" and run.font.size.pt == 9
               for table in document.tables for row in table.rows for cell in row.cells
               for paragraph in cell.paragraphs for run in paragraph.runs if run.text)

    with ZipFile(io.BytesIO(formatted)) as package:
        root = etree.fromstring(package.read("word/document.xml"))
        styles = etree.fromstring(package.read("word/styles.xml"))
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}

        def assert_direct_gadugi(text: str, half_points: str) -> None:
            paragraph = next(node for node in root.xpath(".//w:body//w:p", namespaces=ns)
                             if "".join(node.xpath(".//w:t/text()", namespaces=ns)) == text)
            runs = paragraph.xpath(".//w:r[w:t]", namespaces=ns)
            assert runs
            for run in runs:
                fonts = run.find("w:rPr/w:rFonts", namespaces=ns)
                assert fonts is not None
                assert all(fonts.get(qn(f"w:{name}")) == "Gadugi"
                           for name in ("ascii", "hAnsi", "eastAsia", "cs"))
                assert run.find("w:rPr/w:sz", namespaces=ns).get(qn("w:val")) == half_points
                assert run.find("w:rPr/w:szCs", namespaces=ns).get(qn("w:val")) == half_points

        assert_direct_gadugi("3. HASIL DAN PEMBAHASAN", "20")
        assert_direct_gadugi("Pengujian menggunakan Kp = 2.55 dan menghasilkan akurasi 95%.", "18")
        assert_direct_gadugi("Teks utama menggunakan campuran format dengan Ts = 0.03 s [1].", "20")
        assert_direct_gadugi("Gambar 1: Grafik hasil pengujian", "16")
        assert_direct_gadugi("TABEL I : HASIL PENGUKURAN TEGANGAN", "16")
        for style_id in ("Normal", "Title", "Heading1", "Heading2"):
            style = styles.xpath(f"./w:style[@w:styleId='{style_id}']", namespaces=ns)[0]
            fonts = style.find("w:rPr/w:rFonts", namespaces=ns)
            assert fonts is not None and all(fonts.get(qn(f"w:{name}")) == "Gadugi"
                                             for name in ("ascii", "hAnsi", "eastAsia", "cs"))
        assert len(root.xpath(".//m:oMath|.//m:oMathPara", namespaces=ns)) == 1


def test_caption_detection_is_anchored_and_requires_structural_evidence():
    original = _noncompliant_elkolind_manuscript()
    structure = analyze_manuscript(original)
    kinds = {item.text: item.kind for item in structure.paragraphs}
    assert kinds["Rangkaian pada Gambar 1 menggunakan ESP32 dan sensor ultrasonik."] == "PARAGRAPH"
    assert kinds["Hasil pengukuran ditunjukkan pada Tabel I."] == "PARAGRAPH"
    assert kinds["Gambar 1 : Grafik hasil pengujian"] == "FIGURE_CAPTION"
    assert kinds["Tabel 1: hasil pengukuran tegangan"] == "TABLE_CAPTION"

    document = Document()
    document.add_heading("Judul Artikel", 0)
    document.add_heading("1. PENDAHULUAN", 1)
    uncertain = document.add_paragraph("Gambar 9 : Keterangan tanpa gambar yang berdekatan")
    uncertain.runs[0].font.name = "Times New Roman"
    uncertain.runs[0].font.size = Pt(12)
    uncertain.runs[0].bold = True
    document.add_paragraph("Paragraf biasa setelah kandidat caption.")
    candidate = _save(document)
    profile = extract_style_profile(_elkolind_template())
    rules = {**ELKOLIND_INITIAL_RULES, "left_margin_mm": 14.32, "right_margin_mm": 14.32}
    findings = audit_article(candidate, profile, rules)
    notation = next(item for item in findings if item.check == "Caption notation")
    assert notation.status == "REVIEW_REQUIRED"
    assert all(item.status == "REVIEW_REQUIRED" for item in findings
               if item.target == notation.target and item.category == "Figure")

    selected = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    formatted = apply_safe_fixes(candidate, findings, selected)
    result = Document(io.BytesIO(formatted))
    paragraph = next(item for item in result.paragraphs if item.text.startswith("Gambar 9"))
    assert paragraph.text == "Gambar 9 : Keterangan tanpa gambar yang berdekatan"
    assert paragraph.runs[0].font.name == "Times New Roman"
    assert paragraph.runs[0].font.size.pt == 12
    assert paragraph.runs[0].bold is True


def test_protected_elkolind_captions_are_preserved_and_nonblocking():
    original = _protected_caption_manuscript()
    profile = extract_style_profile(_elkolind_template())
    rules = {**ELKOLIND_INITIAL_RULES, "left_margin_mm": 14.32, "right_margin_mm": 14.32}
    findings = audit_article(original, profile, rules)
    notation = {item.detected: item for item in findings if item.check == "Caption notation"}
    compliant = notation["TABEL II : PERBANDINGAN METODE"]
    warning = notation["Tabel 2: Perbandingan metode"]
    assert compliant.status == "COMPLIANT_PROTECTED"
    assert warning.status == "WARNING_PRESERVED"
    assert "Bookmark" in compliant.protected_object_type and "Field" in compliant.protected_object_type
    assert compliant.preservation_reason and warning.preservation_reason
    assert not any(item.status == "REVIEW_REQUIRED" and item.target in {compliant.target, warning.target}
                   for item in findings)

    selected = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    formatted = apply_safe_fixes(original, findings, selected)

    for text in ("TABEL II : PERBANDINGAN METODE", "Tabel 2: Perbandingan metode"):
        assert _paragraph_xml_by_text(formatted, text) == _paragraph_xml_by_text(original, text)
    assert formatting_integrity(original, formatted)["passed"]


def test_body_inherited_paragraph_wide_bold_is_removed_without_touching_text():
    document = Document()
    document.styles["Normal"].font.bold = True
    document.add_heading("Judul Artikel", 0)
    document.add_heading("1. PENDAHULUAN", 1)
    source_text = "Rangkaian pada Gambar 1 menggunakan ESP32 dan nilai Kp = 2.55."
    document.add_paragraph(source_text)
    original = _save(document)
    profile = extract_style_profile(_elkolind_template())
    rules = {**ELKOLIND_INITIAL_RULES, "left_margin_mm": 14.32, "right_margin_mm": 14.32}
    findings = audit_article(original, profile, rules)
    selected = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    formatted = apply_safe_fixes(original, findings, selected)

    result = Document(io.BytesIO(formatted))
    paragraph = next(item for item in result.paragraphs if item.text == source_text)
    assert paragraph.text == source_text
    assert paragraph.alignment == 3
    assert all(run.font.name == "Gadugi" and run.font.size.pt == 10 and run.bold is False
               for run in paragraph.runs)
    assert result.styles["Normal"].font.bold is False


def test_integrity_blocks_unapproved_numeric_change():
    original = _manuscript()
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original)) as source, ZipFile(output, "w") as destination:
        for member in source.infolist():
            data = source.read(member.filename)
            if member.filename == "word/document.xml":
                assert b"2.55" in data
                data = data.replace(b"2.55", b"1.60")
            destination.writestr(member, data)
    changed = output.getvalue()
    assert not formatting_integrity(original, changed)["passed"]
    assert not formatting_integrity(original, changed)["checks"]["numbers_unchanged"]


def test_safe_web_hyperlinks_allowed_but_external_embeds_rejected():
    from lxml import etree
    original = _manuscript()
    relationships = "http://schemas.openxmlformats.org/package/2006/relationships"

    def with_external(kind: str, target: str) -> bytes:
        output = io.BytesIO()
        with ZipFile(io.BytesIO(original)) as source, ZipFile(output, "w") as destination:
            for member in source.infolist():
                data = source.read(member.filename)
                if member.filename == "word/_rels/document.xml.rels":
                    root = etree.fromstring(data)
                    link = etree.SubElement(root, f"{{{relationships}}}Relationship")
                    link.set("Id", "rId999")
                    link.set("Type", f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}")
                    link.set("Target", target)
                    link.set("TargetMode", "External")
                    data = etree.tostring(root, encoding="UTF-8", xml_declaration=True)
                destination.writestr(member, data)
        return output.getvalue()

    validate_revision_upload("article.docx", with_external("hyperlink", "https://doi.org/10.1234/example"), manuscript=True)
    with pytest.raises(RevisionFileError):
        validate_revision_upload("article.docx", with_external("image", "https://example.com/figure.png"), manuscript=True)
    with pytest.raises(RevisionFileError):
        validate_revision_upload("article.docx", with_external("hyperlink", "file:///C:/private/data"), manuscript=True)


def test_rules_and_report_validation():
    profile = extract_style_profile(_template())
    findings = audit_article(_manuscript(), profile, {"abstract_min_words": 100, "keywords_min": 3})
    assert any(item.category == "Abstract" and item.status == "MISSING_REQUIRED_SECTION" for item in findings)
    assert any(item.category == "Keywords" and item.status == "MISSING_REQUIRED_SECTION" for item in findings)
    content = compliance_report_xlsx(journal="ELKOLIND", manuscript_name="article.docx", template_version=1,
        findings=findings, approved_ids=set(), integrity={"passed": True, "checks": {"numbers_unchanged": True}})
    book = load_workbook(io.BytesIO(content), read_only=True)
    assert book.active["A1"].value == "Template compliance report"
    assert book.active["B2"].value == "ELKOLIND"
    assert book.active["A8"].value == "Category"


def test_journal_template_versioning_and_authorization(case):
    session, store, ja, jb, editor, outsider, finance, submission = case
    first_content = _template()
    first = upload_article_template(session, ja, editor, filename="Article_Elkolind.docx", content=first_content,
        rules={"required_sections": ["Introduction", "References"]}, storage=store)
    second = upload_article_template(session, ja, editor, filename="Article_Elkolind_v2.docx",
        content=_template("Calibri"), storage=store)
    jasens = upload_article_template(session, jb, outsider, filename="Article_Jasens.docx",
        content=_template("Arial"), storage=store)
    assert first.version == jasens.version == 1 and second.version == 2
    assert active_article_template(session, ja.id).id == second.id
    assert active_article_template(session, jb.id).id == jasens.id
    assert len(article_template_versions(session, ja.id)) == 2
    assert authorized_article_template_bytes(session, ja, editor, first.id, storage=store)[1] == first_content
    with pytest.raises(AuthorizationError):
        authorized_article_template_bytes(session, jb, editor, jasens.id, storage=store)
    with pytest.raises(AuthorizationError):
        authorized_article_template_bytes(session, ja, finance, first.id, storage=store)
    activate_article_template(session, ja, editor, first.id)
    assert active_article_template(session, ja.id).id == first.id
    assert second.status.value == "INACTIVE"
    successor = update_article_rules(session, ja, editor, first.id,
        {"abstract_max_words": 250, "keywords_min": 3,
         "style_overrides": {"heading1": {"font": "Arial", "size_pt": 14}}}, storage=store)
    assert successor.version == 3 and active_article_template(session, ja.id).id == successor.id
    assert article_template_config(first)["rules"]["required_sections"] == ["Introduction", "References"]
    assert article_template_config(first)["rules"]["paper"] == "A4"
    assert article_template_config(first)["rules"]["left_margin_mm"] == 20
    assert article_template_config(first)["rules"]["right_margin_mm"] == 20
    assert article_template_config(first)["rules"]["require_lr_margin_confirmation"] is False
    assert article_template_config(successor)["rules"]["style_overrides"]["heading1"]["size_pt"] == 14
    assert first.checksum == __import__("hashlib").sha256(first_content).hexdigest()


def test_complete_template_only_workflow(case):
    session, store, ja, jb, editor, outsider, finance, submission = case
    original = _manuscript_with_abstract_keywords()
    template = upload_article_template(session, ja, editor, filename="Article_Elkolind.docx",
        content=_elkolind_template(), rules={"required_sections": ["Introduction", "References"],
            "left_margin_mm": 14.32, "right_margin_mm": 14.32}, storage=store)
    job = create_article_revision_job(session, ja, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="author_manuscript.docx", content=original,
        metadata={"volume": "13", "issue": "3", "publication_month": "September",
            "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author",
            "received_date": "1 July 2026", "revised_date": "3 August 2026",
            "accepted_date": "10 August 2026"}, storage=store)
    assert job.status == "UPLOADED" and job.template_id == template.id
    findings = audit_article_revision_job(session, ja, editor, job.id, storage=store)
    safe = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    assert safe and job.status == "ANALYZED"
    approve_article_fixes(session, ja, editor, job.id, safe)
    ambiguous = [item for item in findings if item.status == "REVIEW_REQUIRED"]
    assert any(item.category == "Figure" and item.target for item in ambiguous)
    assert job.status == "REVIEW_REQUIRED"
    with pytest.raises(ArticleRevisionRuleError, match="Resolve human review"):
        generate_formatted_manuscript(session, ja, editor, job.id, storage=store)
    resolve_article_review_findings(session, ja, editor, job.id,
        {item.id: "KEEP_AS_IS" for item in ambiguous})
    assert job_resolution_summary(job)["unresolved_human_review"] == 0
    stored_decisions = json.loads(job.approved_fixes_json)["review_decisions"]
    assert stored_decisions and all(item["status"] == "RESOLVED_KEEP"
                                    for item in stored_decisions.values())
    assert job.status == "READY_TO_FORMAT"
    artifacts = generate_formatted_manuscript(session, ja, editor, job.id, storage=store)
    assert len(artifacts) == 2 and job.status == "FORMATTED"
    document = next(item for item in artifacts if item.kind == "FORMATTED_MANUSCRIPT")
    report = next(item for item in artifacts if item.kind == "COMPLIANCE_REPORT")
    filename, formatted = authorized_article_artifact_bytes(session, ja, editor, document.id, storage=store)
    assert filename == "Formatted_Manuscript.docx"
    assert any(paragraph.text == "TABEL I : HASIL PENGUJIAN"
               for paragraph in Document(io.BytesIO(formatted)).paragraphs)
    assert json.loads(job.integrity_json)["passed"]
    with ZipFile(io.BytesIO(formatted)) as output:
        header_text = " ".join(output.read(name).decode("utf-8") for name in output.namelist()
                               if "jasArticle" in name and "header" in name and name.endswith(".xml"))
        footer_text = " ".join(output.read(name).decode("utf-8") for name in output.namelist()
                               if "jasArticle" in name and "footer" in name and name.endswith(".xml"))
        assert "Jurnal Elkolind Volume 13" in header_text
        assert "10.33795/elkolind.v13i3.11397" in header_text
        assert "A. Author: Controlled Revision of a…" in footer_text
        assert "PAGE" in footer_text and "evenAndOddHeaders" in output.read("word/settings.xml").decode()
        assert any("jasArticle" in name and name.startswith("word/media/") for name in output.namelist())
        assert "JASArticle" in output.read("word/styles.xml").decode()
        section_xml = output.read("word/document.xml").decode()
        root = etree.fromstring(section_xml.encode())
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        section = root.xpath(".//w:sectPr", namespaces=ns)[-1]
        reference_order = [
            (etree.QName(node).localname.removesuffix("Reference"),
             node.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type"))
            for node in section
            if etree.QName(node).localname in {"headerReference", "footerReference"}
        ]
        assert reference_order == [
            ("header", "default"), ("footer", "default"), ("header", "even"),
            ("header", "first"), ("footer", "first"), ("footer", "even"),
        ]
        children = [etree.QName(node).localname for node in section]
        assert children.index("titlePg") < children.index("docGrid")
        assert any("http://dx.doi.org/10.33795/elkolind.v13i3.11397" in output.read(name).decode()
                   for name in output.namelist() if "jasArticle" in name and name.endswith(".rels"))
    assert authorized_article_artifact_bytes(session, ja, editor, report.id, storage=store)[0].endswith(".xlsx")
    assert store.read(job.original_storage_key) == original
    with pytest.raises(AuthorizationError):
        authorized_article_artifact_bytes(session, jb, outsider, document.id, storage=store)
    with pytest.raises(AuthorizationError):
        authorized_article_artifact_bytes(session, ja, finance, document.id, storage=store)
    complete_article_revision_job(session, ja, editor, job.id)
    assert job.status == "COMPLETED"
    assert len(article_revision_history(session, ja, editor, job.id)) >= 6


def test_fast_auto_format_generates_with_nonblocking_warnings(case):
    session, store, journal, _other, editor, _outsider, _finance, submission = case
    original = _manuscript_with_abstract_keywords()
    upload_article_template(session, journal, editor, filename="Article_Elkolind.docx",
        content=_elkolind_template(), rules={"required_sections": ["Introduction", "References"],
            "abstract_min_words": 1, "abstract_max_words": 10,
            "left_margin_mm": 14.32, "right_margin_mm": 14.32},
        storage=store)
    job = create_article_revision_job(session, journal, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="author_manuscript.docx", content=original,
        metadata={"volume": "13", "issue": "3", "publication_month": "September",
            "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author"},
        storage=store)

    artifacts = auto_format_and_generate(session, journal, editor, job.id, storage=store)

    assert len(artifacts) == 2
    assert job.status == "FORMATTED"
    result = job_auto_format_result(job)
    assert result and result["auto_fixed"] > 0
    assert result["warnings"] > 0
    assert result["blocking_issues"] == 0
    statuses = {item.status for item in job_fast_findings(job)}
    assert statuses <= {"COMPLIANT", "COMPLIANT_PROTECTED", "AUTO_FIXED", "WARNING",
                        "WARNING_PRESERVED", "REVIEW_REQUIRED", "BLOCKING"}
    assert "BLOCKING" not in statuses
    manuscript = next(item for item in artifacts if item.kind == "FORMATTED_MANUSCRIPT")
    _, formatted = authorized_article_artifact_bytes(session, journal, editor, manuscript.id, storage=store)
    document = Document(io.BytesIO(formatted))
    assert any(paragraph.text == "TABEL I : HASIL PENGUJIAN" for paragraph in document.paragraphs)
    assert json.loads(job.integrity_json)["passed"]
    assert store.read(job.original_storage_key) == original
    actions = {event.action for event in article_revision_history(session, journal, editor, job.id)}
    assert {"ARTICLE_FAST_AUTO_FORMAT_STARTED", "ARTICLE_FAST_AUTO_FORMAT_COMPLETED"} <= actions


def test_fast_generation_preserves_protected_caption_ooxml(case):
    session, store, journal, _other, editor, _outsider, _finance, submission = case
    original = _protected_caption_manuscript()
    upload_article_template(session, journal, editor, filename="Article_Elkolind.docx",
        content=_elkolind_template(), rules={"required_sections": ["Pendahuluan", "Daftar Pustaka"],
            "left_margin_mm": 14.32, "right_margin_mm": 14.32}, storage=store)
    job = create_article_revision_job(session, journal, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="protected_captions.docx", content=original,
        metadata={"volume": "13", "issue": "3", "publication_month": "September",
            "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author"},
        storage=store)

    artifacts = auto_format_and_generate(session, journal, editor, job.id, storage=store)

    result = job_auto_format_result(job)
    assert result["protected_preserved"] == 2
    assert result["blocking_issues"] == 0
    assert job.status == "FORMATTED"
    manuscript = next(item for item in artifacts if item.kind == "FORMATTED_MANUSCRIPT")
    _, formatted = authorized_article_artifact_bytes(session, journal, editor, manuscript.id, storage=store)
    for text in ("TABEL II : PERBANDINGAN METODE", "Tabel 2: Perbandingan metode"):
        assert _paragraph_xml_by_text(formatted, text) == _paragraph_xml_by_text(original, text)
    statuses = {(item.detected, item.status) for item in job_fast_findings(job)
                if item.check == "Caption notation"}
    assert ("TABEL II : PERBANDINGAN METODE", "COMPLIANT_PROTECTED") in statuses
    assert ("Tabel 2: Perbandingan metode", "WARNING_PRESERVED") in statuses

    report = next(item for item in artifacts if item.kind == "COMPLIANCE_REPORT")
    _, report_bytes = authorized_article_artifact_bytes(session, journal, editor, report.id, storage=store)
    sheet = load_workbook(io.BytesIO(report_bytes), read_only=True).active
    headers = [cell.value for cell in sheet[8]]
    assert "Protected-object type" in headers and "Preservation reason" in headers
    rows = list(sheet.iter_rows(min_row=9, values_only=True))
    assert any(row[4] == "COMPLIANT_PROTECTED" and row[5] == "Preserved without modification" for row in rows)
    assert any(row[4] == "WARNING_PRESERVED" and row[5] == "Preserved; warning reported" for row in rows)


def test_fast_mode_retains_batch_accepted_low_risk_reviews(case):
    session, store, journal, _other, editor, _outsider, _finance, submission = case
    upload_article_template(session, journal, editor, filename="Article_Elkolind.docx",
        content=_elkolind_template(), rules={"required_sections": ["Introduction", "References"],
            "left_margin_mm": 14.32, "right_margin_mm": 14.32}, storage=store)
    job = create_article_revision_job(session, journal, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="author_manuscript.docx", content=_manuscript_with_abstract_keywords(),
        metadata={"volume": "13", "issue": "3", "publication_month": "September",
            "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author"},
        storage=store)
    findings = audit_article_revision_job(session, journal, editor, job.id, storage=store)
    assert any(item.status == "REVIEW_REQUIRED" and item.attribute == "caption_text"
               for item in findings)
    accept_all_low_risk_review_findings(session, journal, editor, job.id)
    accepted = {finding_id for finding_id, decision in job_review_decisions(job).items()
                if decision["status"] == "RESOLVED_APPLY_FIX"}
    assert accepted

    artifacts = auto_format_and_generate(session, journal, editor, job.id, storage=store)

    assert len(artifacts) == 2 and job.status == "FORMATTED"
    retained = job_review_decisions(job)
    assert accepted <= set(retained)
    assert all(retained[finding_id]["status"] == "RESOLVED_APPLY_FIX" for finding_id in accepted)


def test_xml_placeholder_scanner_reconstructs_split_word_text_nodes():
    root = etree.fromstring(b"""
        <w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
          <w:p><w:r><w:t>{{public</w:t></w:r><w:r><w:t>ation_month}}</w:t></w:r></w:p>
        </w:hdr>
    """)
    assert "publication_month" in _xml_placeholders(root)


def test_split_run_placeholders_are_detected_inside_shape_xml():
    master = _elkolind_textbox_template(split_runs=True)
    assert elkolind_master_warnings(master) == []
    values = resolve_elkolind_metadata({
        "volume": "13", "issue": "3", "publication_month": "September",
        "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author",
    }, "Controlled Revision of a Novel Method")
    formatted, details = graft_article_master_header_footer(_manuscript(), master, values, elkolind=True)
    assert details["metadata_replacements"] >= 17
    with ZipFile(io.BytesIO(formatted)) as output:
        article_parts = [name for name in output.namelist()
                         if "jasArticle" in name and name.endswith(".xml")]
        xml = b"".join(output.read(name) for name in article_parts)
        assert b"{{" not in xml
        assert b"txbxContent" in xml and b"PAGE" in xml



def test_official_elkolind_doi_suffix_pattern_has_no_dynamic_header_warning():
    warnings = elkolind_master_warnings(_elkolind_textbox_template())
    assert not any("header needs" in warning for warning in warnings)


def test_inactive_even_and_first_header_relationships_are_not_validated():
    document = Document(io.BytesIO(_elkolind_template()))
    document.settings.odd_and_even_pages_header_footer = False
    section = document.sections[0]
    section.different_first_page_header_footer = False
    section.even_page_header.paragraphs[0].text = "Unused even-page header"
    section.first_page_header.paragraphs[0].text = "Unused first-page header"
    inactive_warnings = elkolind_master_warnings(_save(document))
    assert not any("even header needs" in warning for warning in inactive_warnings)
    assert not any("first header needs" in warning for warning in inactive_warnings)

    document.settings.odd_and_even_pages_header_footer = True
    section.different_first_page_header_footer = True
    active_warnings = elkolind_master_warnings(_save(document))
    assert any("even header needs" in warning for warning in active_warnings)
    assert any("first header needs" in warning for warning in active_warnings)


def test_textbox_placeholders_validate_and_generate_without_flattening(case):
    session, store, journal, _other, editor, _outsider, _finance, submission = case
    master = _elkolind_textbox_template()
    assert elkolind_master_warnings(master) == []
    with ZipFile(io.BytesIO(master)) as source:
        source_media = [source.read(name) for name in source.namelist() if name.startswith("word/media/")]
        assert source_media
        source_header_shapes = sum(source.read(name).count(b"txbxContent") for name in source.namelist()
                                   if name.startswith("word/") and "header" in Path(name).name.lower()
                                   and name.endswith(".xml"))
        assert source_header_shapes >= 3
        assert any(b"PAGE" in source.read(name) for name in source.namelist()
                   if name.startswith("word/") and any(kind in Path(name).name.lower()
                                                       for kind in ("header", "footer"))
                   and name.endswith(".xml"))

    template = upload_article_template(session, journal, editor, filename="Official_ELKOLIND_Article.docx",
        content=master, rules={"left_margin_mm": 14.32, "right_margin_mm": 14.32}, storage=store)
    job = create_article_revision_job(session, journal, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="author_manuscript.docx", content=_manuscript_with_abstract_keywords(),
        metadata={"volume": "13", "issue": "3", "publication_month": "September",
            "publication_year": "2026", "doi_suffix": "11397", "first_author": "A. Author"},
        storage=store)
    findings = audit_article_revision_job(session, journal, editor, job.id, storage=store)
    assert any(item.check == "ELKOLIND Article Master header/footer" and item.status == "COMPLIANT"
               for item in findings)
    assert any(item.check == "Apply active Article Master header/footer" and
               item.status == "SAFE_FIX_AVAILABLE" for item in findings)
    assert not any(item.status == "MANUAL_ACTION_REQUIRED" and "header" in item.check.casefold()
                   for item in findings)
    assert not any("Convert the ELKOLIND default header" in item.detected for item in findings)
    approve_article_fixes(session, journal, editor, job.id,
        {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"})
    resolve_article_review_findings(session, journal, editor, job.id,
        {item.id: "KEEP_AS_IS" for item in findings if item.status == "REVIEW_REQUIRED"})
    artifacts = generate_formatted_manuscript(session, journal, editor, job.id, storage=store)
    document = next(item for item in artifacts if item.kind == "FORMATTED_MANUSCRIPT")
    _, formatted = authorized_article_artifact_bytes(session, journal, editor, document.id, storage=store)

    with ZipFile(io.BytesIO(formatted)) as output:
        headers = [name for name in output.namelist()
                   if "jasArticle" in name and "header" in name and name.endswith(".xml")]
        footers = [name for name in output.namelist()
                   if "jasArticle" in name and "footer" in name and name.endswith(".xml")]
        header_xml = b"".join(output.read(name) for name in headers)
        footer_xml = b"".join(output.read(name) for name in footers)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
        header_text = "".join("".join(etree.fromstring(output.read(name)).xpath(".//w:t/text()", namespaces=ns))
                              for name in headers)
        footer_text = "".join("".join(etree.fromstring(output.read(name)).xpath(".//w:t/text()", namespaces=ns))
                              for name in footers)
        assert header_xml.count(b"txbxContent") >= source_header_shapes
        assert b"{{" not in header_xml and b"{{" not in footer_xml
        assert "Jurnal Elkolind Volume 13" in header_text
        assert "10.33795/elkolind.v13i3.11397" in header_text
        assert "A. Author" in footer_text and "Controlled Revision of a" in footer_text
        assert b"PAGE" in header_xml + footer_xml
        copied_media = [output.read(name) for name in output.namelist()
                        if name.startswith("word/media/") and "jasArticle" in name]
        assert copied_media and any(content in source_media for content in copied_media)
        assert b"wps:txbx" in header_xml and b"w:pict" in header_xml
    integrity = json.loads(job.integrity_json)
    assert integrity["passed"]
    assert integrity["checks"]["native_page_field_present"]
    assert integrity["master_header_footer"]["metadata_replacements"] >= 17


def test_elkolind_short_title_and_doi_fallback():
    assert elkolind_master_warnings(_elkolind_template()) == []
    assert short_title_4w("  A   Novel Control   Method for Motors ") == "A Novel Control Method…"
    assert short_title_4w("Short Original Title") == "Short Original Title"
    values = resolve_elkolind_metadata({"volume": "13", "issue": "3", "doi_suffix": "11397"},
        "A Novel Control Method for Motors")
    assert values["doi_full"] == "http://dx.doi.org/10.33795/elkolind.v13i3.11397"
    with pytest.raises(ArticleMetadataError):
        resolve_elkolind_metadata({"volume": "13", "issue": "3", "doi_full": "file:///secret"}, "Title")


def test_elkolind_rule_audit_uses_decided_20mm_left_right_margins():
    findings = audit_article(_manuscript(), extract_style_profile(_elkolind_template()),
        ELKOLIND_INITIAL_RULES)
    assert not any(item.status == "MANUAL_ACTION_REQUIRED" and item.check in {"Left Margin", "Right Margin"}
                   for item in findings)
    margins = [item for item in findings if item.check in {"Left Margin", "Right Margin"}]
    assert len(margins) == 2
    assert all(item.expected == str(__import__("docx").shared.Mm(20).twips) for item in margins)
    assert any(item.check == "Top Margin" and item.expected == str(__import__("docx").shared.Mm(19).twips)
               for item in findings)
    assert any(item.category == "Title" and item.attribute == "size" and item.expected == "48"
               for item in findings)
    assert any(item.category == "References" and item.check == "IEEE numbering" for item in findings)


def test_elkolind_body_dates_are_targeted_placeholders_only():
    manuscript = Document(io.BytesIO(_manuscript()))
    manuscript.add_paragraph("Received {{received_date}}; Revised {{revised_date}}; Accepted {{accepted_date}}")
    original = _save(manuscript)
    values = resolve_elkolind_metadata({"volume": "13", "issue": "3",
        "publication_month": "September", "publication_year": "2026", "doi_suffix": "11397",
        "first_author": "A. Author", "received_date": "1 July 2026",
        "revised_date": "3 August 2026", "accepted_date": "10 August 2026"},
        "Controlled Revision of a Novel Method")
    output, details = graft_article_master_header_footer(original, _elkolind_template(), values, elkolind=True)
    assert details["metadata_replacements"] >= 7
    assert graft_integrity(original, output, values, elkolind=True)["passed"]
    with ZipFile(io.BytesIO(output)) as package:
        body = package.read("word/document.xml").decode()
        assert "Received 1 July 2026; Revised 3 August 2026; Accepted 10 August 2026" in body
        assert "Kp = 2.55" in body


def test_elkolind_metadata_correction_and_template_reselection_are_audited(case):
    session, store, journal, _, editor, _, _, submission = case
    first = upload_article_template(session, journal, editor, filename="ELK_Article_v1.docx",
        content=_elkolind_template(), storage=store)
    job = create_article_revision_job(session, journal, editor, submission=submission,
        title=submission.manuscript_title, submission_identifier="11397", mode="TEMPLATE_ONLY",
        filename="author.docx", content=_manuscript(), metadata={"volume": "13"}, storage=store)
    audit_article_revision_job(session, journal, editor, job.id, storage=store)
    assert job.findings_json
    update_article_job_metadata(session, journal, editor, job.id,
        metadata={"volume": "13", "issue": "3", "first_author": "A. Author"},
        article_title=submission.manuscript_title)
    assert job.status == "UPLOADED" and job.findings_json is None
    second = upload_article_template(session, journal, editor, filename="ELK_Article_v2.docx",
        content=_elkolind_template(), rules={"left_margin_mm": 14.32, "right_margin_mm": 14.32}, storage=store)
    select_active_article_template_for_job(session, journal, editor, job.id)
    assert job.template_id == second.id and first.status.value == "INACTIVE"
    actions = {event.action for event in article_revision_history(session, journal, editor, job.id)}
    assert {"ARTICLE_METADATA_UPDATED", "ARTICLE_TEMPLATE_SELECTED"} <= actions


def test_postgresql_tables_compile_without_changing_existing_schema():
    for table in (ArticleRevisionJob.__table__, ArticleRevisionArtifact.__table__):
        sql = str(CreateTable(table).compile(dialect=postgresql.dialect()))
        assert "CREATE TABLE" in sql and "UUID" in sql
