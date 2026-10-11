"""ELKOLIND front matter: title case, affiliation, abstract table, and submission dates."""

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from services.article_formatting import (
    apply_safe_fixes, audit_article, extract_style_profile, formatting_integrity, title_case_each_word,
)
from services.article_metadata import (
    _replace_tokens, format_metadata_date, resolve_elkolind_metadata, _xml,
)
from services.article_template_service import ELKOLIND_INITIAL_RULES
from services.revision_analysis import NS, analyze_manuscript


def _save(document: Document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _para(document, text, *, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    paragraph.alignment = align
    return paragraph


def _cell_para(cell, text, *, bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, first=False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    paragraph.alignment = align
    return paragraph


def _manuscript() -> bytes:
    """Front matter as authors actually submit it: everything bold/centered, abstract inside a table."""
    document = Document()
    _para(document, "PROTOTIPE PENYIRAMAN OTOMATIS DAN KONTROL IOT UNTUK CABAI ESP32", size=24)
    _para(document, "Muhammad Faiz1, Nina Paramytha2", size=10)
    _para(document, "Jalan Jendral Ahmad Yani No.3, Seberang Ulu I, Kota Palembang,", size=10)
    _para(document, "Sumatera Selatan 30264, Indonesia", size=10)
    _para(document, "*Penulis Korespondensi, e-mail: penulis@example.com", bold=False, size=10)
    table = document.add_table(rows=3, cols=3)
    for cell, label in zip(table.rows[0].cells, ("Received", "Revised", "Accepted")):
        _cell_para(cell, f"{label}:xx/xx/xxxx", first=True)
    top = table.rows[1].cells[0].merge(table.rows[1].cells[2])
    _cell_para(top, "ABSTRAK ", first=True)
    _cell_para(top, "Penyiraman berbasis kondisi media tanam merupakan pendekatan otomasi.")
    _cell_para(top, "Kata Kunci: cabai, penyiraman, Arduino")
    bottom = table.rows[2].cells[0].merge(table.rows[2].cells[2])
    _cell_para(bottom, "ABSTRACT ", first=True)
    _cell_para(bottom, "Irrigation based on growing medium condition is an automation approach.")
    _cell_para(bottom, "Keywords: chili, irrigation, Arduino")
    _para(document, "1. PENDAHULUAN", size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    _para(document, "Penyiraman berbasis kondisi media tanam merupakan salah satu pendekatan.", bold=False,
          align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    return _save(document)


def _template() -> bytes:
    document = Document()
    document.add_paragraph("Template")
    return _save(document)


def _format(manuscript: bytes) -> bytes:
    findings = audit_article(manuscript, extract_style_profile(_template()), dict(ELKOLIND_INITIAL_RULES))
    safe = {item.id for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    assert safe, "expected deterministic ELKOLIND corrections"
    formatted = apply_safe_fixes(manuscript, findings, safe)
    integrity = formatting_integrity(manuscript, formatted)
    assert integrity["passed"], integrity["checks"]
    return formatted


def _cells(document: Document):
    return [paragraph for table in document.tables for row in table.rows for cell in row.cells
            for paragraph in cell.paragraphs if paragraph.text.strip()]


def test_title_case_capitalizes_each_word_and_keeps_acronyms():
    assert title_case_each_word("PROTOTIPE PENYIRAMAN DAN KONTROL IOT UNTUK CABAI") == \
        "Prototipe Penyiraman Dan Kontrol IoT Untuk Cabai"
    assert title_case_each_word("sistem berbasis ESP32 dan pH tanah") == "Sistem Berbasis ESP32 Dan pH Tanah"
    assert title_case_each_word("self-powered (LED) system") == "Self-Powered (LED) System"
    once = title_case_each_word("PROTOTIPE IOT LDR")
    assert title_case_each_word(once) == once


def test_front_matter_is_not_classified_as_authors():
    structure = analyze_manuscript(_manuscript())
    assert structure.authors == ("Muhammad Faiz1, Nina Paramytha2", "*Penulis Korespondensi, e-mail: penulis@example.com")
    assert len(structure.affiliations) == 2 and structure.affiliations[1].startswith("Sumatera Selatan 30264")


def test_title_affiliation_dates_and_abstract_follow_elkolind_template():
    formatted = _format(_manuscript())
    document = Document(io.BytesIO(formatted))
    body = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    title, authors, aff1, aff2 = body[0], body[1], body[2], body[3]
    assert title.text == "Prototipe Penyiraman Otomatis Dan Kontrol IoT Untuk Cabai ESP32"
    assert all(run.font.size == Pt(9) and run.bold is False for run in aff1.runs + aff2.runs)
    assert aff1.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert all(run.bold for run in authors.runs)

    cells = {paragraph.text.split(":")[0].strip(): paragraph for paragraph in _cells(document)}
    for label in ("Received", "Revised", "Accepted"):
        assert all(run.bold is False and run.font.size == Pt(10) for run in cells[label].runs)
    for heading in ("ABSTRAK", "ABSTRACT"):
        paragraph = cells[heading]
        assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        assert all(run.bold and run.font.size == Pt(9) for run in paragraph.runs)
    abstract = next(p for p in _cells(document) if p.text.startswith("Penyiraman berbasis"))
    assert abstract.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert all(run.bold is False and run.font.size == Pt(9) for run in abstract.runs)
    for label in ("Kata Kunci", "Keywords"):
        keywords = cells[label]
        assert keywords.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
        bold = [run.text for run in keywords.runs if run.bold]
        plain = [run.text for run in keywords.runs if not run.bold]
        assert "".join(bold).strip().rstrip(":") == label and plain and all(run.font.size == Pt(9) for run in keywords.runs)


def test_formatting_is_idempotent_once_compliant():
    formatted = _format(_manuscript())
    findings = audit_article(formatted, extract_style_profile(_template()), dict(ELKOLIND_INITIAL_RULES))
    assert not [item for item in findings if item.status == "SAFE_FIX_AVAILABLE"]


def _set_numpr(paragraph, *, num_id: str, ilvl: int) -> None:
    """Attach Word's own multilevel-list numbering (no literal "1." typed into the run text).

    This is how Word's Format > Bullets and Numbering (multilevel list) feature renders
    section numbers: the "1."/"2.1" prefix is drawn from word/numbering.xml at render time and
    never appears in the paragraph's own w:t text, which is exactly what analyze_manuscript
    must recognize a heading from without a real numbering.xml (it only reads w:numPr/w:numId
    and w:ilvl off the paragraph itself).
    """
    from docx.oxml.ns import qn

    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.makeelement(qn("w:numPr"), {})
    ilvl_el = p_pr.makeelement(qn("w:ilvl"), {qn("w:val"): str(ilvl)})
    num_id_el = p_pr.makeelement(qn("w:numId"), {qn("w:val"): num_id})
    num_pr.append(ilvl_el)
    num_pr.append(num_id_el)
    p_pr.insert(0, num_pr)


def _manuscript_with_word_numbered_headings() -> bytes:
    """Section headings numbered by Word's list feature, like job 11454's real manuscript.

    Regression for a bug where a manuscript using this (common) numbering style was never
    recognized past its front matter at all: _kind()'s heading regex needs a literal leading
    digit in the paragraph text, which a Word-numbered heading does not have (the paragraph
    text is just "PENDAHULUAN", not "1. PENDAHULUAN"), and falls back to a "Heading N" Word
    style, which a numbered-list paragraph ("List Paragraph" style) also doesn't carry. With
    neither signal, every paragraph for the rest of the document -- the entire body, not just
    the real front matter -- stayed classified as FRONT_MATTER, so every body paragraph's exact
    text got appended to structure.authors/.affiliations and (via _role() in
    article_formatting.py) every body paragraph was then formatted as a bold, centered author
    line instead of normal justified body text: "Integrity" wasn't the problem, the whole
    article just came out looking wrong, which is what the user saw and reported.
    """
    document = Document()
    _para(document, "RANCANG BANGUN ALAT UJI ADC0804 DAN DAC0808-LM741", size=24)
    _para(document, "Fauziah Sholikhatun Nisa", size=10)
    _para(document, "Jurusan Teknik Elektro, Politeknik Negeri Malang", size=10)
    heading = _para(document, "PENDAHULUAN", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    _set_numpr(heading, num_id="1", ilvl=0)
    intro = document.add_paragraph()
    intro.add_run("Perkembangan teknologi elektronika digital mendorong penggunaan sistem.")
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    subheading = _para(document, "Tahapan Penelitian", bold=True, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
    _set_numpr(subheading, num_id="1", ilvl=1)
    method = document.add_paragraph()
    method.add_run("Penelitian ini merancang dan mengkarakterisasi alat uji.")
    method.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return _save(document)


def test_word_numbered_section_headings_close_the_front_matter():
    structure = analyze_manuscript(_manuscript_with_word_numbered_headings())
    heading = next(item for item in structure.paragraphs if item.text.strip() == "PENDAHULUAN")
    assert heading.kind == "SECTION_HEADING"
    assert heading.section_id == "SECTION_1"
    subheading = next(item for item in structure.paragraphs if item.text.strip() == "Tahapan Penelitian")
    assert subheading.kind == "SECTION_HEADING"
    assert subheading.section_id == "SECTION_1_1"
    body = next(item for item in structure.paragraphs
                if item.text.startswith("Perkembangan teknologi elektronika digital"))
    assert body.section_id != "FRONT_MATTER"
    assert body.text not in structure.authors
    assert body.text not in structure.affiliations


def test_submission_dates_replace_xx_placeholders():
    assert format_metadata_date("2026-07-01") == "01/07/2026"
    assert format_metadata_date("03/08/2026") == "03/08/2026"
    assert format_metadata_date("10 August 2026") == "10 August 2026"
    formatted = _format(_manuscript())
    values = resolve_elkolind_metadata({"received_date": "2026-07-01", "revised_date": "03/08/2026",
                                        "accepted_date": ""}, "Judul")
    with __import__("zipfile").ZipFile(io.BytesIO(formatted)) as archive:
        root = _xml(archive.read("word/document.xml"))
    _replace_tokens(root, values)
    text = " ".join(root.xpath(".//w:t/text()", namespaces=NS))
    assert "Received:01/07/2026" in text and "Revised:03/08/2026" in text
    assert "Accepted:xx/xx/xxxx" in text  # no date supplied: placeholder left for the editor
