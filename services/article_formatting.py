"""Deterministic article-template audit and conservative OOXML formatting patches."""

from __future__ import annotations

import copy
import io
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Mm
from lxml import etree
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from services.revision_analysis import (
    FIGURE_CAPTION_RE, NS, TABLE_CAPTION_RE, W, analyze_manuscript, citation_tokens, numeric_tokens,
)
from services.revision_docx import ApprovedPatch, check_integrity


class ArticleFormattingError(ValueError):
    pass


@dataclass(frozen=True)
class Finding:
    id: str
    category: str
    check: str
    expected: str
    detected: str
    status: str
    target: str | None = None
    attribute: str | None = None
    value: object = None
    protected_object_type: str | None = None
    preservation_reason: str | None = None


PAGE_KEYS = ("page_width", "page_height", "top_margin", "bottom_margin", "left_margin", "right_margin")
STYLE_KEYS = ("font", "size", "bold", "alignment", "space_before", "space_after", "line_spacing", "first_line_indent")
ROLE_STYLE = {"title": "Title", "body": "Normal", "heading1": "Heading 1", "heading2": "Heading 2",
              "caption": "Caption", "author": "Author", "affiliation": "Affiliation", "abstract": "Abstract",
              "keywords": "Keywords", "reference": "Reference", "figure_caption": "Figure Caption",
              "table_caption": "Table Caption"}
W_TAG = f"{{{W}}}"
ELKOLIND_FORMAT_TARGET = "ELKOLIND_SEMANTIC_FORMATTING"
ELKOLIND_FORMAT_FIX_ID = "document:elkolind_semantic_formatting:apply"
ELKOLIND_SECTION_NAMES = {
    "pendahuluan", "introduction", "metode penelitian", "methodology", "methods",
    "hasil dan pembahasan", "results and discussion", "kesimpulan", "conclusion",
    "ucapan terimakasih", "ucapan terima kasih", "acknowledgements", "acknowledgments",
    "daftar pustaka", "references", "bibliography",
}
ELKOLIND_ROLE_SPECS: dict[str, dict[str, object]] = {
    "title": {"font": "Gadugi", "size": 48, "alignment": 1},
    "author": {"font": "Gadugi", "size": 20, "alignment": 1, "bold": True},
    "corresponding_author": {"font": "Gadugi", "size": 20, "alignment": 1},
    "affiliation": {"font": "Gadugi", "size": 18, "alignment": 1, "bold": False},
    "date": {"font": "Gadugi", "size": 20, "alignment": 1, "bold": False},
    "abstract_heading": {"font": "Gadugi", "size": 18, "alignment": 3, "bold": True},
    "abstract": {"font": "Gadugi", "size": 18, "alignment": 3, "bold": False},
    "keywords": {"font": "Gadugi", "size": 18},
    # Keywords inside the abstract table: only the "Kata Kunci:" / "Keywords:" label is bold (see _keyword_label_ok).
    "keywords_table": {"font": "Gadugi", "size": 18, "alignment": 3},
    "body": {"font": "Gadugi", "size": 20, "alignment": 3},
    "heading1": {"font": "Gadugi", "size": 20, "alignment": 0, "bold": True},
    "heading2": {"font": "Gadugi", "size": 20, "alignment": 0, "bold": True},
    "figure_caption": {"font": "Gadugi", "size": 16, "alignment": 1},
    "table_caption": {"font": "Gadugi", "size": 16, "alignment": 1},
    "table_cell": {"font": "Gadugi", "size": 18},
    "reference": {"font": "Gadugi", "size": 16},
}
ELKOLIND_STYLE_SPECS = {
    "normal": {**ELKOLIND_ROLE_SPECS["body"], "bold": False},
    "title": ELKOLIND_ROLE_SPECS["title"],
    "heading1": ELKOLIND_ROLE_SPECS["heading1"],
    "heading2": ELKOLIND_ROLE_SPECS["heading2"],
    "author": ELKOLIND_ROLE_SPECS["author"],
    "affiliation": ELKOLIND_ROLE_SPECS["affiliation"],
    "abstract": ELKOLIND_ROLE_SPECS["abstract"],
    "keywords": {**ELKOLIND_ROLE_SPECS["keywords"], "bold": False},
    "reference": ELKOLIND_ROLE_SPECS["reference"],
    "caption": {**ELKOLIND_ROLE_SPECS["figure_caption"], "bold": False},
    "figurecaption": {**ELKOLIND_ROLE_SPECS["figure_caption"], "bold": False},
    "tablecaption": {**ELKOLIND_ROLE_SPECS["table_caption"], "bold": False},
}


def _parse(data: bytes) -> etree._Element:
    return etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def _style_spec(style) -> dict[str, object]:
    result: dict[str, object] = {}
    for item in (style, *list(_base_styles(style))):
        if "font" not in result and item.font.name:
            result["font"] = item.font.name
        if "size" not in result and item.font.size:
            result["size"] = int(round(item.font.size.pt * 2))
        if "bold" not in result and item.font.bold is not None:
            result["bold"] = bool(item.font.bold)
        fmt = item.paragraph_format
        for name, value in (("alignment", fmt.alignment), ("space_before", fmt.space_before),
                            ("space_after", fmt.space_after), ("first_line_indent", fmt.first_line_indent)):
            if name not in result and value is not None:
                result[name] = int(value) if name == "alignment" else value.twips
        line = item.element.xpath("./w:pPr/w:spacing/@w:line")
        if "line_spacing" not in result and line:
            rule = item.element.xpath("./w:pPr/w:spacing/@w:lineRule")
            result["line_spacing"] = {"line": int(line[0]), "rule": rule[0] if rule else None}
    return result


def _base_styles(style):
    seen = set()
    current = style.base_style
    while current is not None and current.style_id not in seen:
        seen.add(current.style_id)
        yield current
        current = current.base_style


def extract_style_profile(content: bytes) -> dict[str, object]:
    """Use explicit properties only; theme/default ambiguity remains for editor review."""
    document = Document(io.BytesIO(content))
    section = document.sections[0]
    page = {key: getattr(section, key).twips for key in PAGE_KEYS}
    styles = {}
    for role, name in ROLE_STYLE.items():
        if name in document.styles:
            styles[role] = _style_spec(document.styles[name])
    with ZipFile(io.BytesIO(content)) as archive:
        root = _parse(archive.read("word/document.xml"))
        columns = root.xpath(".//w:sectPr/w:cols/@w:num", namespaces=NS)
        footer_xml = b"".join(archive.read(name) for name in archive.namelist()
                              if name.startswith("word/footer") and name.endswith(".xml"))
    structure = analyze_manuscript(content)
    return {"page": page, "orientation": int(section.orientation), "styles": styles, "sections": len(document.sections),
            "columns": int(columns[0]) if columns else 1,
            "has_page_number": b"PAGE" in footer_xml,
            "has_abstract": any(item.kind == "ABSTRACT_HEADING" for item in structure.paragraphs),
            "has_keywords": bool(structure.keywords)}


def _effective_style(paragraph: etree._Element, document) -> dict[str, object]:
    style_ids = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
    style = next((item for item in document.styles if style_ids and item.style_id == style_ids[0]),
                 document.styles["Normal"])
    result = _style_spec(style)
    for key, query in (("alignment", "./w:pPr/w:jc/@w:val"),
                       ("space_before", "./w:pPr/w:spacing/@w:before"),
                       ("space_after", "./w:pPr/w:spacing/@w:after"),
                       ("first_line_indent", "./w:pPr/w:ind/@w:firstLine")):
        values = paragraph.xpath(query, namespaces=NS)
        if values:
            if key == "alignment":
                result[key] = {"left": 0, "center": 1, "right": 2, "both": 3, "justify": 3}.get(values[0], values[0])
            else:
                result[key] = int(values[0])
    line = paragraph.xpath("./w:pPr/w:spacing/@w:line", namespaces=NS)
    if line:
        rule = paragraph.xpath("./w:pPr/w:spacing/@w:lineRule", namespaces=NS)
        result["line_spacing"] = {"line": int(line[0]), "rule": rule[0] if rule else None}
    runs = paragraph.xpath("./w:r[w:t]", namespaces=NS)
    if runs:
        run = runs[0]
        font = run.xpath("./w:rPr/w:rFonts/@w:ascii", namespaces=NS)
        size = run.xpath("./w:rPr/w:sz/@w:val", namespaces=NS)
        if font:
            result["font"] = font[0]
        if size:
            result["size"] = int(size[0])
    return result


def _simple_text_paragraph(paragraph: etree._Element) -> bool:
    if not paragraph.xpath("./w:r/w:t", namespaces=NS):
        return False
    for child in paragraph:
        if child.tag == W_TAG + "pPr":
            continue
        if child.tag != W_TAG + "r":
            return False
        if any(node.tag not in {W_TAG + "rPr", W_TAG + "t"} for node in child):
            return False
    return True


def _protected_word_object_types(paragraph: etree._Element) -> tuple[str, ...]:
    """Describe OOXML constructs that make in-place caption rewriting unsafe."""
    checks = (
        ("Bookmark", ".//w:bookmarkStart|.//w:bookmarkEnd"),
        ("Field", ".//w:fldSimple|.//w:fldChar|.//w:instrText"),
        ("Hyperlink", ".//w:hyperlink"),
        ("Content control", ".//w:sdt|.//w:sdtContent"),
        ("Comment anchor", ".//w:commentRangeStart|.//w:commentRangeEnd|.//w:commentReference"),
        ("Footnote/endnote reference", ".//w:footnoteReference|.//w:endnoteReference"),
        ("Drawing/object", ".//w:drawing|.//w:pict|.//w:object"),
    )
    found = [label for label, query in checks if paragraph.xpath(query, namespaces=NS)]
    if not found and not _simple_text_paragraph(paragraph):
        found.append("Other protected OOXML")
    return tuple(found)


def _protected_caption_status(paragraph: etree._Element, info, *, confirmed: bool,
                              normalized: str | None) -> tuple[str, str, str] | None:
    protected = _protected_word_object_types(paragraph)
    if not protected:
        return None
    object_type = ", ".join(protected)
    compliant = confirmed and normalized is not None and normalized == info.text
    status = "COMPLIANT_PROTECTED" if compliant else "WARNING_PRESERVED"
    reason = (
        "Protected Word objects detected; caption already follows ELKOLIND notation and was preserved without modification."
        if compliant else
        "Protected Word objects detected; automatic caption or formatting changes were skipped to preserve the original OOXML."
    )
    return status, object_type, reason


def _ordinary_text_runs(paragraph: etree._Element) -> list[etree._Element]:
    """Return visible Word text runs while excluding Office Math objects."""
    return [run for run in paragraph.xpath(".//w:r[w:t]", namespaces=NS)
            if not any(etree.QName(ancestor).namespace == NS["m"] for ancestor in run.iterancestors())]


def _on_off_value(element: etree._Element | None) -> bool:
    if element is None:
        return False
    return (element.get(W_TAG + "val") or "true").lower() not in {"0", "false", "off", "no"}


def _direct_run_format_matches(paragraph: etree._Element, attribute: str, expected: object) -> bool:
    runs = _ordinary_text_runs(paragraph)
    if not runs:
        return False
    for run in runs:
        properties = run.find(W_TAG + "rPr")
        if properties is None:
            return False
        if attribute == "font":
            fonts = properties.find(W_TAG + "rFonts")
            if fonts is None or any(fonts.get(W_TAG + name) != str(expected)
                                    for name in ("ascii", "hAnsi", "eastAsia", "cs")):
                return False
        elif attribute == "size":
            if any((properties.find(W_TAG + name) is None or
                    properties.find(W_TAG + name).get(W_TAG + "val") != str(expected))
                   for name in ("sz", "szCs")):
                return False
        elif attribute == "bold":
            if _on_off_value(properties.find(W_TAG + "b")) != bool(expected):
                return False
        else:
            return False
    return True


def _roman_number(value: int) -> str:
    if not 1 <= value <= 3999:
        raise ArticleFormattingError("Table caption number is outside the supported Roman numeral range.")
    result = []
    for number, numeral in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"),
                            (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
                            (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        count, value = divmod(value, number)
        result.append(numeral * count)
    return "".join(result)


def _roman_value(value: str) -> int | None:
    if value.isdigit():
        number = int(value)
        return number if 1 <= number <= 3999 else None
    values = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100, "D": 500, "M": 1000}
    total, prior = 0, 0
    for character in reversed(value.upper()):
        current = values.get(character)
        if current is None:
            return None
        total += -current if current < prior else current
        prior = max(prior, current)
    return total if total and _roman_number(total) == value.upper() else None


def _normalized_caption(kind: str, text: str) -> str | None:
    if kind == "FIGURE_CAPTION":
        match = FIGURE_CAPTION_RE.fullmatch(text)
        if not match:
            return None
        suffix = match.group(2).strip()
        return f"Gambar {int(match.group(1))}: {suffix}"
    if kind == "TABLE_CAPTION":
        match = TABLE_CAPTION_RE.fullmatch(text)
        if not match:
            return None
        number = _roman_value(match.group(1))
        if number is None:
            return None
        suffix = match.group(2).strip().upper()
        return f"TABEL {_roman_number(number)} : {suffix}"
    return None


def _caption_position_confirmed(paragraph: etree._Element, kind: str) -> bool:
    """Require the journal's object/caption order before applying a caption fix."""
    body = paragraph.getparent()
    if body is None or body.tag != W_TAG + "body":
        return False
    siblings = list(body)
    position = siblings.index(paragraph)
    direction = -1 if kind == "FIGURE_CAPTION" else 1
    neighbour = position + direction
    while 0 <= neighbour < len(siblings):
        candidate = siblings[neighbour]
        if candidate.tag == W_TAG + "p" and not candidate.xpath(
            ".//w:t/text()|.//w:drawing|.//w:pict", namespaces=NS
        ):
            neighbour += direction
            continue
        break
    if not 0 <= neighbour < len(siblings):
        return False
    candidate = siblings[neighbour]
    if kind == "TABLE_CAPTION":
        return candidate.tag == W_TAG + "tbl"
    return bool(candidate.xpath(".//w:drawing|.//w:pict", namespaces=NS))


def _style_bold(styles: etree._Element | None, style_id: str | None) -> bool:
    if styles is None:
        return False
    seen: set[str] = set()
    current_id = style_id
    while current_id and current_id not in seen:
        seen.add(current_id)
        matches = styles.xpath("./w:style[@w:styleId=$style_id]", namespaces=NS, style_id=current_id)
        if not matches:
            break
        style = matches[0]
        bold = style.find(W_TAG + "rPr/" + W_TAG + "b")
        if bold is not None:
            return _on_off_value(bold)
        based_on = style.find(W_TAG + "basedOn")
        current_id = based_on.get(W_TAG + "val") if based_on is not None else None
    default_bold = styles.find(W_TAG + "docDefaults/" + W_TAG + "rPrDefault/" + W_TAG + "rPr/" + W_TAG + "b")
    return _on_off_value(default_bold)


def _paragraph_all_runs_bold(paragraph: etree._Element, styles: etree._Element | None = None) -> bool:
    runs = _ordinary_text_runs(paragraph)
    if not runs:
        return False
    paragraph_style = paragraph.find(W_TAG + "pPr/" + W_TAG + "pStyle")
    paragraph_style_id = paragraph_style.get(W_TAG + "val") if paragraph_style is not None else None
    if paragraph_style_id is None and styles is not None:
        defaults = styles.xpath(
            "./w:style[@w:type='paragraph' and @w:default='1']/@w:styleId",
            namespaces=NS,
        )
        paragraph_style_id = defaults[0] if defaults else None
    paragraph_bold = _style_bold(styles, paragraph_style_id)
    effective: list[bool] = []
    for run in runs:
        direct = run.find(W_TAG + "rPr/" + W_TAG + "b")
        if direct is not None:
            effective.append(_on_off_value(direct))
            continue
        run_style = run.find(W_TAG + "rPr/" + W_TAG + "rStyle")
        if run_style is not None:
            effective.append(_style_bold(styles, run_style.get(W_TAG + "val")))
        else:
            effective.append(paragraph_bold)
    return all(effective)


def _replace_paragraph_text(paragraph: etree._Element, replacement: str) -> None:
    nodes = paragraph.xpath("./w:r/w:t", namespaces=NS)
    if not nodes:
        raise ArticleFormattingError("A caption selected for normalization has no plain text runs.")
    cursor = 0
    for index, node in enumerate(nodes):
        if index == len(nodes) - 1:
            value = replacement[cursor:]
        else:
            length = len(node.text or "")
            value = replacement[cursor:cursor + length]
            cursor += length
        node.text = value
        if value.startswith(" ") or value.endswith(" "):
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        else:
            node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)


DATE_LABEL_RE = re.compile(r"^\s*(?:received|revised|accepted)\s*:", re.I)
ABSTRACT_LABEL_RE = re.compile(r"^\s*(?:abstrak|abstract)\b", re.I)
KEYWORDS_LABEL_RE = re.compile(r"^\s*(?:keywords|kata kunci)\s*:", re.I)
TITLE_ACRONYMS = {
    "IOT": "IoT", "WIFI": "WiFi", "LED": "LED", "LDR": "LDR", "PLC": "PLC", "PWM": "PWM", "ADC": "ADC",
    "GPS": "GPS", "GSM": "GSM", "LCD": "LCD", "USB": "USB", "RFID": "RFID", "MQTT": "MQTT", "API": "API",
    "AI": "AI", "ML": "ML", "CNN": "CNN", "LSTM": "LSTM", "IEEE": "IEEE", "UV": "UV", "PID": "PID",
    "MPPT": "MPPT", "PV": "PV", "IP": "IP", "AC": "AC", "DC": "DC", "SMS": "SMS", "CCTV": "CCTV",
    "DHT": "DHT", "ESP": "ESP", "PCB": "PCB", "MPU": "MPU", "IMU": "IMU", "RTC": "RTC", "HMI": "HMI",
}


def title_case_each_word(text: str) -> str:
    """Capitalize the first letter of every word; keep acronyms and mixed-case tokens (IoT, ESP32, pH)."""
    def fix(word: str) -> str:
        match = re.fullmatch(r"([^A-Za-z0-9]*)(.*?)([^A-Za-z0-9]*)", word, re.S)
        lead, core, trail = match.groups() if match else ("", word, "")
        if not core:
            return word
        if "-" in core:
            return lead + "-".join(fix(part) for part in core.split("-")) + trail
        letters = re.sub(r"[^A-Za-z]", "", core)
        if not letters or any(ch.isdigit() for ch in core):
            return word
        if core.upper() in TITLE_ACRONYMS:
            return lead + TITLE_ACRONYMS[core.upper()] + trail
        if core.islower() or core.isupper() or core == core.capitalize():
            return lead + core[:1].upper() + core[1:].lower() + trail
        return word
    return re.sub(r"\S+", lambda m: fix(m.group(0)), text)


def _front_table_role(info, structure) -> str | None:
    """Classify front-matter table cells: submission dates and the ABSTRAK / ABSTRACT block."""
    if info.kind != "TABLE_CELL" or info.section_id != "FRONT_MATTER":
        return None
    in_abstract = False
    for item in structure.paragraphs:
        if item.section_id != "FRONT_MATTER":
            break
        text = item.text.strip()
        if item.kind != "TABLE_CELL":
            if text:
                in_abstract = False
            continue
        role = None
        if DATE_LABEL_RE.match(text):
            role, in_abstract = "date", False
        elif ABSTRACT_LABEL_RE.match(text) and len(text) < 40:
            role, in_abstract = "abstract_heading", True
        elif KEYWORDS_LABEL_RE.match(text):
            role, in_abstract = "keywords_table", False
        elif in_abstract and text:
            role = "abstract"
        if item.index == info.index:
            return role
    return None


def _keyword_label_length(paragraph: etree._Element) -> int:
    text = "".join(paragraph.xpath(".//w:r[w:t]/w:t/text()", namespaces=NS))
    match = KEYWORDS_LABEL_RE.match(text)
    return match.end() if match else 0


def _keyword_label_ok(paragraph: etree._Element) -> bool:
    """True when only the leading 'Keywords:' label is bold."""
    label_end = _keyword_label_length(paragraph)
    position = 0
    for run in _ordinary_text_runs(paragraph):
        length = len("".join(run.xpath("./w:t/text()", namespaces=NS)))
        properties = run.find(W_TAG + "rPr")
        bold = _on_off_value(properties.find(W_TAG + "b")) if properties is not None else False
        if length and bold != (position < label_end):
            return False
        position += length
    return True


def _format_keywords_table(paragraph: etree._Element, spec: dict[str, object]) -> None:
    """Apply the keyword spec, then bold only the label (splitting a run at the colon if needed)."""
    label_end = _keyword_label_length(paragraph)
    position = 0
    for run in list(_ordinary_text_runs(paragraph)):
        texts = run.xpath("./w:t", namespaces=NS)
        run_text = "".join(node.text or "" for node in texts)
        length = len(run_text)
        if length and position < label_end < position + length and len(texts) == 1:
            split_at = label_end - position
            tail = copy.deepcopy(run)
            tail_node = tail.xpath("./w:t", namespaces=NS)[0]
            texts[0].text, tail_node.text = run_text[:split_at], run_text[split_at:]
            for node in (texts[0], tail_node):
                if (node.text or "").startswith(" ") or (node.text or "").endswith(" "):
                    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.addnext(tail)
            _set_run_format(run, {**spec, "bold": True})
            _set_run_format(tail, {**spec, "bold": False})
        else:
            _set_run_format(run, {**spec, "bold": position < label_end})
        position += length


def _role(info, first_text_index: int, structure) -> str | None:
    if not info.text.strip():
        return None
    if info.index == first_text_index:
        return "title"
    front_table = _front_table_role(info, structure)
    if front_table:
        return front_table
    if info.kind == "ABSTRACT_HEADING":
        return "abstract_heading"
    if info.kind == "KEYWORDS":
        return "keywords"
    if info.kind == "PARAGRAPH" and info.section_id == "ABSTRACT":
        return "abstract"
    if info.text in structure.references:
        return "reference"
    if info.section_id == "FRONT_MATTER" and info.text in structure.affiliations:
        return "affiliation"
    if info.section_id == "FRONT_MATTER" and info.text in structure.authors:
        if re.search(r"correspond|e-?mail|@|\*", info.text, re.I):
            return "corresponding_author"
        return "author"
    if info.kind == "SECTION_HEADING":
        return "heading2" if re.match(r"^\d+\.\d+", info.text.strip()) else "heading1"
    normalized = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", info.text.strip()).casefold()
    if normalized in ELKOLIND_SECTION_NAMES:
        return "heading2" if re.match(r"^\d+\.\d+", info.text.strip()) else "heading1"
    if info.kind == "FIGURE_CAPTION":
        return "figure_caption"
    if info.kind == "TABLE_CAPTION":
        return "table_caption"
    if info.kind == "TABLE_CELL":
        return "table_cell"
    if info.kind == "PARAGRAPH" and info.section_id != "FRONT_MATTER":
        return "body"
    return None


def _finding(category: str, check: str, expected: object, detected: object, status: str,
             *, target: str | None = None, attribute: str | None = None,
             protected_object_type: str | None = None,
             preservation_reason: str | None = None) -> Finding:
    identifier = f"{category.lower().replace(' ', '_')}:{target or 'document'}:{attribute or check.lower().replace(' ', '_')}"
    return Finding(identifier, category, check, str(expected), str(detected), status, target, attribute, expected,
                   protected_object_type, preservation_reason)


def _elkolind_spec(role: str, text: str) -> dict[str, object] | None:
    base = ELKOLIND_ROLE_SPECS.get(role)
    if base is None:
        return None
    result = dict(base)
    if role == "figure_caption":
        result["alignment"] = 3 if "\n" in text or len(text) > 60 else 1
    return result


def _paragraph_format_matches(paragraph: etree._Element, document: Document,
                              spec: dict[str, object]) -> bool:
    actual = _effective_style(paragraph, document)
    for attribute, expected in spec.items():
        if attribute in {"font", "size", "bold"}:
            if not _direct_run_format_matches(paragraph, attribute, expected):
                return False
        elif actual.get(attribute) != expected:
            return False
    return True


def _style_format_matches(style: etree._Element, spec: dict[str, object]) -> bool:
    run = style.find(W_TAG + "rPr")
    paragraph = style.find(W_TAG + "pPr")
    if "font" in spec:
        fonts = run.find(W_TAG + "rFonts") if run is not None else None
        if fonts is None or any(fonts.get(W_TAG + name) != spec["font"]
                                for name in ("ascii", "hAnsi", "eastAsia", "cs")):
            return False
    if "size" in spec:
        if run is None or any(run.find(W_TAG + name) is None or
                              run.find(W_TAG + name).get(W_TAG + "val") != str(spec["size"])
                              for name in ("sz", "szCs")):
            return False
    if "bold" in spec:
        actual_bold = _on_off_value(run.find(W_TAG + "b")) if run is not None else False
        if actual_bold != bool(spec["bold"]):
            return False
    if "alignment" in spec:
        alignment = paragraph.find(W_TAG + "jc") if paragraph is not None else None
        expected = {0: "left", 1: "center", 2: "right", 3: "both"}[int(spec["alignment"])]
        if alignment is None or alignment.get(W_TAG + "val") != expected:
            return False
    return True


def _elkolind_semantic_compliant(root: etree._Element, styles: etree._Element,
                                 document: Document, structure) -> bool:
    paragraphs = root.xpath(".//w:body//w:p", namespaces=NS)
    first_text = next((item.index for item in structure.paragraphs if item.text.strip()), -1)
    for info in structure.paragraphs:
        role = _role(info, first_text, structure)
        spec = _elkolind_spec(role or "", info.text)
        if spec is None:
            continue
        paragraph = paragraphs[info.index - 1]
        if role in {"figure_caption", "table_caption"} and not _caption_position_confirmed(paragraph, info.kind):
            # Caption-like prose without the expected adjacent object remains an
            # editor-review item and is intentionally excluded from auto-formatting.
            continue
        if role in {"figure_caption", "table_caption"} and _protected_word_object_types(paragraph):
            # Protected caption OOXML is reported separately and intentionally
            # excluded from the semantic auto-fix aggregate.
            continue
        if not _paragraph_format_matches(paragraph, document, spec):
            return False
        if role == "keywords_table" and not _keyword_label_ok(paragraph):
            return False
        if role == "title" and _simple_text_paragraph(paragraph) and title_case_each_word(info.text) != info.text:
            return False
        if role in {"body", "keywords", "figure_caption", "table_caption"} and _paragraph_all_runs_bold(paragraph, styles):
            return False
        normalized = _normalized_caption(info.kind, info.text)
        if normalized is not None and normalized != info.text:
            return False
    for style in styles.xpath("./w:style", namespaces=NS):
        style_id = (style.get(W_TAG + "styleId") or "").replace(" ", "").casefold()
        spec = ELKOLIND_STYLE_SPECS.get(style_id)
        if spec is not None and not _style_format_matches(style, spec):
            return False
    return True


def audit_article(manuscript: bytes, profile: dict[str, object], rules: dict[str, object]) -> list[Finding]:
    structure = analyze_manuscript(manuscript)
    document = Document(io.BytesIO(manuscript))
    with ZipFile(io.BytesIO(manuscript)) as archive:
        root = _parse(archive.read("word/document.xml"))
        styles_root = _parse(archive.read("word/styles.xml"))
        names = set(archive.namelist())
        footer_xml = b"".join(archive.read(name) for name in names
                              if name.startswith("word/footer") and name.endswith(".xml"))
    paragraphs = root.xpath(".//w:body//w:p", namespaces=NS)
    first_text = next((item.index for item in structure.paragraphs if item.text.strip()), -1)
    findings: list[Finding] = []
    target_page = dict(profile.get("page", {}))
    if rules.get("paper") == "A4":
        target_page.update(page_width=Mm(210).twips, page_height=Mm(297).twips)
    elif rules.get("paper") == "LETTER":
        target_page.update(page_width=Mm(215.9).twips, page_height=Mm(279.4).twips)
    for edge in ("top", "bottom", "left", "right"):
        if rules.get(f"{edge}_margin_mm") is not None:
            target_page[f"{edge}_margin"] = Mm(float(rules[f"{edge}_margin_mm"])).twips
    single_section = len(document.sections) == 1 and profile.get("sections") == 1
    for key in PAGE_KEYS:
        expected = target_page.get(key)
        if expected is None:
            continue
        detected = getattr(document.sections[0], key).twips
        status = "COMPLIANT" if expected == detected else ("SAFE_FIX_AVAILABLE" if single_section else "REVIEW_REQUIRED")
        findings.append(_finding("Document", key.replace("_", " ").title(), expected, detected,
                                 status, target="section", attribute=key))
    if profile.get("orientation") is not None:
        expected = profile["orientation"]
        detected = int(document.sections[0].orientation)
        findings.append(_finding("Document", "Orientation", expected, detected,
            "COMPLIANT" if expected == detected else ("SAFE_FIX_AVAILABLE" if single_section else "REVIEW_REQUIRED"),
            target="section", attribute="orientation"))
    columns = root.xpath(".//w:sectPr/w:cols/@w:num", namespaces=NS)
    actual_columns = int(columns[0]) if columns else 1
    expected_columns = (1 if rules.get("layout") == "single_column" else
                        2 if rules.get("layout") == "two_columns" else profile.get("columns", 1))
    findings.append(_finding("Document", "Columns", expected_columns, actual_columns,
                             "COMPLIANT" if actual_columns == expected_columns else
                             ("SAFE_FIX_AVAILABLE" if single_section else "REVIEW_REQUIRED"),
                             target="section", attribute="columns"))
    findings.append(_finding("Document", "Section count", profile.get("sections", 1), len(document.sections),
                             "COMPLIANT" if profile.get("sections", 1) == len(document.sections) else "REVIEW_REQUIRED"))
    for part in ("header", "footer"):
        expected_has = bool(profile.get(f"has_{part}"))
        actual_has = any(name.startswith(f"word/{part}") and name.endswith(".xml") for name in names)
        if expected_has:
            findings.append(_finding("Document", part.title(), "Present", "Present" if actual_has else "Missing",
                                     "REVIEW_REQUIRED" if actual_has else "MANUAL_ACTION_REQUIRED"))
    if profile.get("has_page_number"):
        numbered = b"PAGE" in footer_xml
        findings.append(_finding("Document", "Page numbering", "Present", "Present" if numbered else "Missing",
                                 "REVIEW_REQUIRED" if numbered else "MANUAL_ACTION_REQUIRED"))
    style_profile = {role: dict(spec) for role, spec in profile.get("styles", {}).items()}
    for role in ("figure_caption", "table_caption"):
        if "caption" in style_profile:
            style_profile[role] = dict(style_profile["caption"])
    if rules.get("default_font"):
        style_profile.setdefault("body", {})["font"] = rules["default_font"]
    if rules.get("body_font"):
        style_profile.setdefault("body", {})["font"] = rules["body_font"]
    if rules.get("body_font_size"):
        style_profile.setdefault("body", {})["size"] = int(round(float(rules["body_font_size"]) * 2))
    for role, prefix in (("title", "title"), ("body", "body"), ("abstract", "abstract")):
        if rules.get(f"{prefix}_font"):
            style_profile.setdefault(role, {})["font"] = rules[f"{prefix}_font"]
        if rules.get(f"{prefix}_size_pt"):
            style_profile.setdefault(role, {})["size"] = int(round(float(rules[f"{prefix}_size_pt"]) * 2))
    if rules.get("title_alignment"):
        style_profile.setdefault("title", {})["alignment"] = {
            "left": 0, "center": 1, "right": 2, "justify": 3}[rules["title_alignment"]]
    for role, rule_name in (("author", "author_format"), ("affiliation", "affiliation_format"),
                            ("abstract", "abstract_format"), ("keywords", "keyword_format"),
                            ("body", "body_format"), ("reference", "reference_format"),
                            ("figure_caption", "figure_caption_format"),
                            ("table_caption", "table_caption_format")):
        spec = rules.get(rule_name, {})
        for name in ("font", "size_pt", "alignment"):
            if name not in spec:
                continue
            value = spec[name]
            attribute = "size" if name == "size_pt" else name
            style_profile.setdefault(role, {})[attribute] = (
                int(round(float(value) * 2)) if name == "size_pt" else
                {"left": 0, "center": 1, "right": 2, "justify": 3}.get(value, value)
            )
    for role, overrides in rules.get("style_overrides", {}).items():
        target = style_profile.setdefault(role, {})
        for name, value in overrides.items():
            if name == "size_pt":
                target["size"] = int(round(value * 2))
            elif name in {"space_before_pt", "space_after_pt", "first_line_indent_pt"}:
                target[name.removesuffix("_pt")] = int(round(value * 20))
            elif name == "line_spacing_multiplier":
                target["line_spacing"] = {"line": int(round(value * 240)), "rule": "auto"}
            else:
                target[name] = value
    elkolind_profile = rules.get("formatting_profile") == "ELKOLIND"
    if elkolind_profile:
        for role, spec in ELKOLIND_ROLE_SPECS.items():
            style_profile[role] = dict(spec)
        semantic_compliant = _elkolind_semantic_compliant(root, styles_root, document, structure)
        findings.append(Finding(
            id=ELKOLIND_FORMAT_FIX_ID,
            category="Document",
            check="ELKOLIND deterministic typography and caption formatting",
            expected="Gadugi semantic typography with normalized figure and table captions",
            detected="Compliant" if semantic_compliant else "Author formatting overrides remain",
            status="COMPLIANT" if semantic_compliant else "SAFE_FIX_AVAILABLE",
            target=ELKOLIND_FORMAT_TARGET,
            attribute="elkolind_semantic_format",
            value="ELKOLIND",
        ))
        role_records = [(item, _role(item, first_text, structure)) for item in structure.paragraphs]

        def summary_status(roles: set[str], *, caption: bool = False) -> tuple[str, str]:
            selected = [(item, role) for item, role in role_records if role in roles]
            if not selected:
                return "REVIEW_REQUIRED", "No confidently classified paragraph"
            compliant = True
            checked = 0
            ambiguous = False
            protected_warnings = 0
            protected_compliant = 0
            for item, role in selected:
                paragraph = paragraphs[item.index - 1]
                if role in {"figure_caption", "table_caption"} and not _caption_position_confirmed(paragraph, item.kind):
                    if caption:
                        ambiguous = True
                    continue
                checked += 1
                if role in {"figure_caption", "table_caption"}:
                    protected = _protected_caption_status(
                        paragraph, item, confirmed=True,
                        normalized=_normalized_caption(item.kind, item.text),
                    )
                    if protected:
                        if protected[0] == "COMPLIANT_PROTECTED":
                            protected_compliant += 1
                        else:
                            protected_warnings += 1
                        continue
                spec = _elkolind_spec(role or "", item.text) or {}
                if not _paragraph_format_matches(paragraph, document, spec):
                    compliant = False
                if role in {"body", "keywords", "figure_caption", "table_caption"} and _paragraph_all_runs_bold(paragraph, styles_root):
                    compliant = False
                if caption and _normalized_caption(item.kind, item.text) != item.text:
                    compliant = False
            if not checked:
                return "REVIEW_REQUIRED", "No structurally confirmed caption paragraph"
            if ambiguous:
                return "REVIEW_REQUIRED", "At least one caption-like paragraph lacks adjacent object evidence"
            if protected_warnings:
                return "WARNING_PRESERVED", f"{protected_warnings} protected caption(s) preserved without modification"
            if protected_compliant and protected_compliant == checked:
                return "COMPLIANT_PROTECTED", f"{protected_compliant} protected caption(s) already compliant and preserved"
            return (("COMPLIANT", f"{checked} paragraph(s) compliant") if compliant else
                    ("SAFE_FIX_AVAILABLE", "Deterministic ELKOLIND formatting correction available"))

        all_roles = {role for role in ELKOLIND_ROLE_SPECS}
        summary_checks = (
            ("Global font", all_roles, False),
            ("Abstract font", {"abstract", "abstract_heading"}, False),
            ("Abstract body bold", {"abstract"}, False),
            ("Body font", {"body"}, False),
            ("Body paragraph-wide bold", {"body"}, False),
            ("Heading 1", {"heading1"}, False),
            ("Heading 2", {"heading2"}, False),
            ("Figure caption format", {"figure_caption"}, True),
            ("Table caption format", {"table_caption"}, True),
        )
        for check, roles, caption in summary_checks:
            status, detected = summary_status(roles, caption=caption)
            findings.append(Finding(
                id=f"elkolind_summary:{check.lower().replace(' ', '_')}",
                category="ELKOLIND",
                check=check,
                expected="Compliant with deterministic ELKOLIND formatting rules",
                detected=detected,
                status=status,
                target=ELKOLIND_FORMAT_TARGET,
                attribute="elkolind_semantic_format" if status == "SAFE_FIX_AVAILABLE" else None,
                value="ELKOLIND" if status == "SAFE_FIX_AVAILABLE" else None,
            ))
    for info in structure.paragraphs:
        role = _role(info, first_text, structure)
        if role is None or role not in style_profile:
            continue
        paragraph = paragraphs[info.index - 1]
        if role != "table_cell" and any(ancestor.tag == W_TAG + "tbl" for ancestor in paragraph.iterancestors()):
            continue
        if role == "title" and not any(value.lower() == "title" for value in paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)):
            findings.append(_finding("Title", "Title location", "Explicit Word Title style", "Uncertain first paragraph", "REVIEW_REQUIRED"))
            continue
        actual = _effective_style(paragraph, document)
        simple = _simple_text_paragraph(paragraph)
        expected_style = (_elkolind_spec(role, info.text) if elkolind_profile else None) or dict(style_profile[role])
        caption_confirmed = (
            role not in {"figure_caption", "table_caption"}
            or _caption_position_confirmed(paragraph, info.kind)
        )
        protected_caption = (
            _protected_caption_status(
                paragraph, info, confirmed=caption_confirmed,
                normalized=_normalized_caption(info.kind, info.text),
            )
            if role in {"figure_caption", "table_caption"} else None
        )
        if role == "figure_caption" and rules.get("figure_caption_format"):
            caption_rule = rules["figure_caption_format"]
            if "\n" in info.text:
                alignment = caption_rule.get("multi_line_alignment")
            elif len(info.text) <= 60:
                alignment = caption_rule.get("single_line_alignment")
            else:
                alignment = None
            if alignment:
                expected_style["alignment"] = {"left": 0, "center": 1, "right": 2, "justify": 3}[alignment]
        if role == "figure_caption" and len(info.text) > 60 and "\n" not in info.text:
            # Word line wrapping depends on the uploaded figure width and fonts.
            expected_style.pop("alignment", None)
            findings.append(_finding("Figure", "Caption line alignment", "Center if one line; justify if multiple",
                "Rendered line count unknown", "REVIEW_REQUIRED", target=f"PARAGRAPH_{info.index:03d}"))
        for attribute in STYLE_KEYS:
            expected = expected_style.get(attribute)
            if expected is None:
                continue
            direct_attribute = attribute in {"font", "size", "bold"}
            compliant = (_direct_run_format_matches(paragraph, attribute, expected)
                         if direct_attribute else actual.get(attribute) == expected)
            detected = (actual.get(attribute) if not direct_attribute else
                        expected if compliant else "Mixed, inherited, or conflicting direct formatting")
            safely_patchable = bool(_ordinary_text_runs(paragraph)) if direct_attribute else simple
            if protected_caption:
                status = "COMPLIANT_PROTECTED" if compliant and protected_caption[0] == "COMPLIANT_PROTECTED" else "WARNING_PRESERVED"
                patch_attribute = None
            else:
                status = (
                    "REVIEW_REQUIRED" if not caption_confirmed else
                    "COMPLIANT" if compliant else
                    "SAFE_FIX_AVAILABLE" if safely_patchable else "REVIEW_REQUIRED"
                )
                patch_attribute = attribute if status != "REVIEW_REQUIRED" or safely_patchable else None
            findings.append(_finding(role.title(), f"{role.title()} {attribute.replace('_', ' ')}", expected,
                                     detected if detected is not None else "Unknown", status,
                                     target=f"PARAGRAPH_{info.index:03d}", attribute=patch_attribute,
                                     protected_object_type=protected_caption[1] if protected_caption else None,
                                     preservation_reason=protected_caption[2] if protected_caption else None))
    if rules.get("title_case") == "sentence_case":
        findings.append(_finding("Title", "Sentence case", "Sentence case while preserving acronyms",
            structure.title, "REVIEW_REQUIRED"))
    section_names = {re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", p.text).strip().casefold()
                     for p in structure.paragraphs if p.kind == "SECTION_HEADING"}
    for section in rules.get("required_sections", []):
        found = section.strip().casefold() in section_names
        findings.append(_finding("Structure", f"Required section: {section}", "Present", "Present" if found else "Missing",
                                 "COMPLIANT" if found else "MISSING_REQUIRED_SECTION"))
    if profile.get("has_abstract"):
        findings.append(_finding("Abstract", "Abstract present", "Present", "Present" if structure.abstract else "Missing",
                                 "COMPLIANT" if structure.abstract else "MISSING_REQUIRED_SECTION"))
    if profile.get("has_keywords"):
        findings.append(_finding("Keywords", "Keywords present", "Present", "Present" if structure.keywords else "Missing",
                                 "COMPLIANT" if structure.keywords else "MISSING_REQUIRED_SECTION"))
    abstract_words = len(structure.abstract.split())
    if rules.get("abstract_min_words") is not None or rules.get("abstract_max_words") is not None:
        minimum, maximum = rules.get("abstract_min_words"), rules.get("abstract_max_words")
        valid = bool(structure.abstract) and (minimum is None or abstract_words >= minimum) and (maximum is None or abstract_words <= maximum)
        findings.append(_finding("Abstract", "Abstract word count", f"{minimum or 0}–{maximum or 'no maximum'} words",
                                 abstract_words, "COMPLIANT" if valid else ("MISSING_REQUIRED_SECTION" if not structure.abstract else "REVIEW_REQUIRED")))
    keywords = [word.strip() for word in re.split(r"[;,]", re.sub(r"^(keywords|kata kunci)\s*:\s*", "", structure.keywords, flags=re.I)) if word.strip()]
    if rules.get("keywords_min") is not None or rules.get("keywords_max") is not None:
        minimum, maximum = rules.get("keywords_min"), rules.get("keywords_max")
        valid = bool(keywords) and (minimum is None or len(keywords) >= minimum) and (maximum is None or len(keywords) <= maximum)
        findings.append(_finding("Keywords", "Keyword count", f"{minimum or 0}–{maximum or 'no maximum'}",
                                 len(keywords), "COMPLIANT" if valid else ("MISSING_REQUIRED_SECTION" if not keywords else "REVIEW_REQUIRED")))
    if rules.get("title_max_words") is not None:
        count = len(structure.title.split())
        findings.append(_finding("Title", "Title word count", f"≤ {rules['title_max_words']}", count,
                                 "COMPLIANT" if count <= rules["title_max_words"] else "REVIEW_REQUIRED"))
    if elkolind_profile:
        for info in structure.paragraphs:
            normalized = _normalized_caption(info.kind, info.text)
            if normalized is None:
                continue
            paragraph = paragraphs[info.index - 1]
            confirmed = _caption_position_confirmed(paragraph, info.kind)
            protected = _protected_caption_status(paragraph, info, confirmed=confirmed, normalized=normalized)
            status = (protected[0] if protected else
                      "REVIEW_REQUIRED" if not confirmed else
                      "COMPLIANT" if normalized == info.text else
                      "SAFE_FIX_AVAILABLE" if _simple_text_paragraph(paragraph) else "REVIEW_REQUIRED")
            findings.append(_finding(
                "Figure" if info.kind == "FIGURE_CAPTION" else "Table",
                "Caption notation", normalized, info.text, status,
                target=f"PARAGRAPH_{info.index:03d}",
                attribute=None if protected else "caption_text",
                protected_object_type=protected[1] if protected else None,
                preservation_reason=protected[2] if protected else None,
            ))
    else:
        for category, captions, prefix in (("Figure", structure.figure_captions, rules.get("figure_caption_prefix")),
                                            ("Table", structure.table_captions, rules.get("table_caption_prefix"))):
            if prefix:
                for index, caption in enumerate(captions, 1):
                    findings.append(_finding(category, f"Caption {index} prefix", prefix, caption[:len(prefix)],
                                             "COMPLIANT" if caption.casefold().startswith(prefix.casefold()) else "REVIEW_REQUIRED"))
    body = root.find(W_TAG + "body")
    body_children = list(body) if body is not None else []
    for info in structure.paragraphs:
        if info.kind not in {"FIGURE_CAPTION", "TABLE_CAPTION"}:
            continue
        paragraph = paragraphs[info.index - 1]
        if paragraph.getparent() is not body:
            continue
        position = body_children.index(paragraph)
        if info.kind == "FIGURE_CAPTION":
            spec = rules.get("figure_caption_format", {})
            object_kind, direction, category = "drawing", -1, "Figure"
            number_format = rules.get("figure_format", {}).get("numbering")
            numbered = bool(re.match(r"^(?:Figure|Fig\.?|Gambar)\s+\d+\b", info.text, re.I))
        else:
            spec = rules.get("table_caption_format", {})
            object_kind, direction, category = "table", 1, "Table"
            number_format = rules.get("table_format", {}).get("numbering")
            numbered = bool(re.match(r"^(?:Table|Tabel)\s+[IVXLCDM]+\b", info.text))
        if number_format and not elkolind_profile:
            findings.append(_finding(category, "Caption numbering", number_format,
                info.text.split(".", 1)[0][:45], "COMPLIANT" if numbered else "REVIEW_REQUIRED",
                target=f"PARAGRAPH_{info.index:03d}"))
        if spec.get("position"):
            neighbour = position + direction
            while 0 <= neighbour < len(body_children):
                candidate = body_children[neighbour]
                if candidate.tag == W_TAG + "p" and not candidate.xpath(".//w:t/text()|.//w:drawing|.//w:pict", namespaces=NS):
                    neighbour += direction
                    continue
                break
            adjacent = 0 <= neighbour < len(body_children) and (
                body_children[neighbour].tag == W_TAG + "tbl" if object_kind == "table" else
                bool(body_children[neighbour].xpath(".//w:drawing|.//w:pict", namespaces=NS))
            )
            findings.append(_finding(category, "Caption placement", spec["position"],
                "Adjacent to object" if adjacent else "Not adjacent / object uncertain",
                "COMPLIANT" if adjacent else "REVIEW_REQUIRED", target=f"PARAGRAPH_{info.index:03d}"))
        if category == "Table" and spec.get("case") == "template" and not elkolind_profile:
            findings.append(_finding("Table", "Caption capitalization", "Match official master wording",
                info.text, "REVIEW_REQUIRED", target=f"PARAGRAPH_{info.index:03d}"))
    if structure.authors:
        findings.append(_finding("Authors", "Author order and affiliation markers", "Match submission metadata and template",
                                 f"{len(structure.authors)} front-matter line(s)", "REVIEW_REQUIRED"))
    else:
        findings.append(_finding("Authors", "Author block", "Present", "Not reliably identified", "MANUAL_ACTION_REQUIRED"))
    findings.append(_finding("Affiliations", "Affiliation text and markers", "Match template and verified authors",
                             f"{len(structure.affiliations)} front-matter line(s)", "REVIEW_REQUIRED"))
    findings.append(_finding("Headings", "Hierarchy and numbering", "Follow article master",
                             f"{len(section_names)} identified section(s)", "REVIEW_REQUIRED"))
    for category, count, captions in (("Figure", structure.figure_count, structure.figure_captions),
                                       ("Table", structure.table_count, structure.table_captions)):
        if count:
            findings.append(_finding(category, "Caption count and placement", f"At least {count} caption(s); inspect placement",
                                     len(captions), "REVIEW_REQUIRED" if len(captions) >= count else "MANUAL_ACTION_REQUIRED"))
            findings.append(_finding(category, "Object alignment", "Follow article master",
                                     f"{count} existing object(s)", "REVIEW_REQUIRED"))
    if structure.equation_count:
        findings.append(_finding("Equation", "Numbering and alignment", "Follow article master",
                                 f"{structure.equation_count} preserved Word equation(s)", "REVIEW_REQUIRED"))
        if rules.get("equation_format", {}).get("numbering"):
            numbers = [int(value) for value in re.findall(r"\((\d+)\)", " ".join(
                info.text for info in structure.paragraphs if info.kind == "PARAGRAPH" and
                info.text.strip().startswith("(")))]
            findings.append(_finding("Equation", "Arabic parenthesized numbering", "(1), (2), …",
                str(numbers) if numbers else "No reliably identified equation labels", "REVIEW_REQUIRED"))
    if rules.get("reference_style") and structure.references:
        findings.append(_finding("References", "Citation and reference style", rules["reference_style"],
                                 f"{len(structure.references)} preserved reference entries", "REVIEW_REQUIRED"))
    elif structure.references:
        findings.append(_finding("References", "Citation and reference formatting", "Match article master",
                                 f"{len(structure.references)} preserved reference entries", "REVIEW_REQUIRED"))
    if rules.get("reference_format", {}).get("numbering") == "bracketed_arabic" and structure.references:
        labels = [re.match(r"^\[(\d+)\]", entry.strip()) for entry in structure.references]
        valid = all(match is not None and int(match.group(1)) == position
                    for position, match in enumerate(labels, 1))
        findings.append(_finding("References", "IEEE numbering", "[1], [2], …",
            "Sequential" if valid else "Missing or non-sequential labels",
            "COMPLIANT" if valid else "REVIEW_REQUIRED"))
    if not structure.references and "references" in {section.casefold() for section in rules.get("required_sections", [])}:
        findings.append(_finding("References", "Reference entries", "Present", "Missing", "MANUAL_ACTION_REQUIRED"))
    return findings


def compliance_score(findings: list[Finding]) -> int:
    measured = [item for item in findings if item.status != "REVIEW_REQUIRED"]
    if not measured:
        return 0
    compliant = {"COMPLIANT", "COMPLIANT_PROTECTED", "AUTO_FIXED"}
    return round(100 * sum(item.status in compliant for item in measured) / len(measured))


def _ensure(parent: etree._Element, tag: str, *, first: bool = False) -> etree._Element:
    child = parent.find(W_TAG + tag)
    if child is None:
        child = etree.Element(W_TAG + tag)
        parent.insert(0, child) if first else parent.append(child)
    return child


def _set_run_format(run: etree._Element, spec: dict[str, object]) -> None:
    properties = _ensure(run, "rPr", first=True)
    if "font" in spec:
        fonts = _ensure(properties, "rFonts", first=True)
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(W_TAG + name, str(spec["font"]))
    if "size" in spec:
        for name in ("sz", "szCs"):
            _ensure(properties, name).set(W_TAG + "val", str(spec["size"]))
    if "bold" in spec:
        _ensure(properties, "b").set(W_TAG + "val", "1" if bool(spec["bold"]) else "0")


def _set_paragraph_alignment(paragraph: etree._Element, alignment: object) -> None:
    name = {0: "left", 1: "center", 2: "right", 3: "both"}.get(int(alignment))
    if name is None:
        raise ArticleFormattingError("Unsupported paragraph alignment in ELKOLIND formatting profile.")
    _ensure(_ensure(paragraph, "pPr", first=True), "jc").set(W_TAG + "val", name)


def _set_style_format(style: etree._Element, spec: dict[str, object]) -> None:
    properties = _ensure(style, "rPr")
    if "font" in spec:
        fonts = _ensure(properties, "rFonts", first=True)
        for name in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(W_TAG + name, str(spec["font"]))
    if "size" in spec:
        for name in ("sz", "szCs"):
            _ensure(properties, name).set(W_TAG + "val", str(spec["size"]))
    if "bold" in spec:
        _ensure(properties, "b").set(W_TAG + "val", "1" if bool(spec["bold"]) else "0")
    if "alignment" in spec:
        paragraph = _ensure(style, "pPr")
        name = {0: "left", 1: "center", 2: "right", 3: "both"}[int(spec["alignment"])]
        _ensure(paragraph, "jc").set(W_TAG + "val", name)


def _apply_elkolind_semantic_formatting(root: etree._Element, styles: etree._Element, structure) -> None:
    paragraphs = root.xpath(".//w:body//w:p", namespaces=NS)
    first_text = next((item.index for item in structure.paragraphs if item.text.strip()), -1)
    for info in structure.paragraphs:
        role = _role(info, first_text, structure)
        spec = _elkolind_spec(role or "", info.text)
        if spec is None:
            continue
        paragraph = paragraphs[info.index - 1]
        if role in {"figure_caption", "table_caption"} and not _caption_position_confirmed(paragraph, info.kind):
            continue
        if role in {"figure_caption", "table_caption"} and _protected_word_object_types(paragraph):
            # Bookmarks, fields, hyperlinks, content controls and other complex
            # OOXML remain byte-for-byte untouched inside the caption paragraph.
            continue
        run_spec = dict(spec)
        if role in {"body", "keywords", "figure_caption", "table_caption"} and _paragraph_all_runs_bold(paragraph, styles):
            # A whole paragraph in bold is treated as accidental source
            # formatting for these semantic roles. Local emphasis in mixed
            # paragraphs is preserved.
            run_spec["bold"] = False
        normalized = _normalized_caption(info.kind, info.text)
        if normalized is not None and normalized != info.text:
            _replace_paragraph_text(paragraph, normalized)
        if role == "title" and _simple_text_paragraph(paragraph):
            titled = title_case_each_word(info.text)
            if titled != info.text:
                _replace_paragraph_text(paragraph, titled)
        if role == "keywords_table":
            _format_keywords_table(paragraph, run_spec)
        else:
            for run in _ordinary_text_runs(paragraph):
                _set_run_format(run, run_spec)
        if "alignment" in spec:
            _set_paragraph_alignment(paragraph, spec["alignment"])
    for style in styles.xpath("./w:style", namespaces=NS):
        style_id = (style.get(W_TAG + "styleId") or "").replace(" ", "").casefold()
        spec = ELKOLIND_STYLE_SPECS.get(style_id)
        if spec is not None:
            _set_style_format(style, spec)


def _apply_paragraph_value(paragraph: etree._Element, attribute: str, value: object) -> None:
    if attribute in {"font", "size", "bold"}:
        spec = {attribute: value}
        for run in _ordinary_text_runs(paragraph):
            _set_run_format(run, spec)
        return
    properties = _ensure(paragraph, "pPr", first=True)
    if attribute == "alignment":
        _set_paragraph_alignment(paragraph, value)
    elif attribute == "caption_text":
        _replace_paragraph_text(paragraph, str(value))
    elif attribute in {"space_before", "space_after"}:
        _ensure(properties, "spacing").set(W_TAG + ("before" if attribute == "space_before" else "after"), str(value))
    elif attribute == "first_line_indent":
        _ensure(properties, "ind").set(W_TAG + "firstLine", str(value))
    elif attribute == "line_spacing":
        if not isinstance(value, dict) or "line" not in value:
            raise ArticleFormattingError("Invalid line-spacing profile.")
        spacing = _ensure(properties, "spacing")
        spacing.set(W_TAG + "line", str(value["line"]))
        if value.get("rule"):
            spacing.set(W_TAG + "lineRule", str(value["rule"]))
        else:
            spacing.attrib.pop(W_TAG + "lineRule", None)
    else:
        raise ArticleFormattingError("Unsupported safe formatting attribute.")


def apply_safe_fixes(original: bytes, findings: list[Finding], approved_ids: set[str]) -> bytes:
    safe = {item.id: item for item in findings if item.status == "SAFE_FIX_AVAILABLE"}
    if not approved_ids <= set(safe):
        raise ArticleFormattingError("Only findings classified as safe may be applied.")
    structure = analyze_manuscript(original)
    with ZipFile(io.BytesIO(original)) as source:
        root = _parse(source.read("word/document.xml"))
        styles = _parse(source.read("word/styles.xml"))
        paragraphs = root.xpath(".//w:body//w:p", namespaces=NS)
        sections = root.xpath(".//w:sectPr", namespaces=NS)
        semantic_ids = {item.id for item in safe.values()
                        if item.target == ELKOLIND_FORMAT_TARGET and
                        item.attribute == "elkolind_semantic_format"}
        if approved_ids & semantic_ids:
            _apply_elkolind_semantic_formatting(root, styles, structure)
        for finding_id in sorted(approved_ids):
            item = safe[finding_id]
            if item.target == ELKOLIND_FORMAT_TARGET and item.attribute == "elkolind_semantic_format":
                continue
            if item.target == "section":
                if len(sections) != 1 or item.attribute not in (*PAGE_KEYS, "orientation", "columns"):
                    raise ArticleFormattingError("Section layout is too complex for a safe automatic fix.")
                section = sections[0]
                if item.attribute == "columns":
                    _ensure(section, "cols").set(W_TAG + "num", str(item.value))
                elif item.attribute == "orientation":
                    _ensure(section, "pgSz").set(W_TAG + "orient", "landscape" if int(item.value) else "portrait")
                elif item.attribute in {"page_width", "page_height"}:
                    element = _ensure(section, "pgSz")
                    element.set(W_TAG + ("w" if item.attribute == "page_width" else "h"), str(item.value))
                else:
                    element = _ensure(section, "pgMar")
                    element.set(W_TAG + item.attribute.split("_")[0], str(item.value))
            else:
                match = re.fullmatch(r"PARAGRAPH_(\d+)", item.target or "")
                if not match or int(match.group(1)) > len(paragraphs):
                    raise ArticleFormattingError("Formatting target is no longer present.")
                paragraph = paragraphs[int(match.group(1)) - 1]
                direct_only = item.attribute in {"font", "size", "bold"}
                if direct_only and not _ordinary_text_runs(paragraph):
                    raise ArticleFormattingError("Formatting target has no ordinary Word text runs.")
                if not direct_only and not _simple_text_paragraph(paragraph):
                    raise ArticleFormattingError("Formatting target contains protected Word objects.")
                _apply_paragraph_value(paragraph, item.attribute or "", item.value)
        output = io.BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as destination:
            for member in source.infolist():
                if member.filename == "word/document.xml":
                    data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
                elif member.filename == "word/styles.xml":
                    data = etree.tostring(styles, encoding="UTF-8", xml_declaration=True, standalone=True)
                else:
                    data = source.read(member.filename)
                destination.writestr(member, data)
    return output.getvalue()


def _caption_integrity_patches(original: bytes, formatted: bytes) -> list[ApprovedPatch]:
    before = analyze_manuscript(original)
    after = analyze_manuscript(formatted)
    revised = {item.identifier: item for item in after.paragraphs}
    patches = []
    first_text = next((item for item in before.paragraphs if item.text.strip()), None)
    if first_text is not None:
        titled = title_case_each_word(first_text.text)
        current_title = revised.get(first_text.identifier)
        if titled != first_text.text and current_title is not None and current_title.text == titled:
            patches.append(ApprovedPatch(first_text.identifier, first_text.text, titled, None))
    for item in before.paragraphs:
        if item.kind not in {"FIGURE_CAPTION", "TABLE_CAPTION"}:
            continue
        expected = _normalized_caption(item.kind, item.text)
        current = revised.get(item.identifier)
        if expected is not None and expected != item.text and current is not None and current.text == expected:
            reason = ("ELKOLIND table-caption numeral normalization"
                      if numeric_tokens(item.text) != numeric_tokens(expected) else None)
            patches.append(ApprovedPatch(item.identifier, item.text, expected, reason))
    return patches


def formatting_integrity(original: bytes, formatted: bytes) -> dict[str, object]:
    caption_patches = _caption_integrity_patches(original, formatted)
    legacy = check_integrity(original, formatted, caption_patches)
    with ZipFile(io.BytesIO(original)) as before, ZipFile(io.BytesIO(formatted)) as after:
        names_before, names_after = set(before.namelist()), set(after.namelist())
        original_xml, final_xml = _parse(before.read("word/document.xml")), _parse(after.read("word/document.xml"))
        allowed_parts = {"word/document.xml", "word/styles.xml"}
        unchanged_parts = all(before.read(name) == after.read(name) for name in names_before - allowed_parts) if names_before == names_after else False
    paths = {"numbers": ".//w:t", "equations": ".//m:oMath|.//m:oMathPara", "figures": ".//w:drawing|.//w:pict",
             "tables": ".//w:tbl", "captions": ".//w:p"}
    texts_equal = not legacy.checks["unexpected_text_changes"]
    equation_xml_equal = [etree.tostring(n) for n in original_xml.xpath(paths["equations"], namespaces=NS)] == [etree.tostring(n) for n in final_xml.xpath(paths["equations"], namespaces=NS)]
    figure_xml_equal = [etree.tostring(n) for n in original_xml.xpath(paths["figures"], namespaces=NS)] == [etree.tostring(n) for n in final_xml.xpath(paths["figures"], namespaces=NS)]
    original_structure = analyze_manuscript(original)
    final_structure = analyze_manuscript(formatted)
    protected_before = " ".join(item.text for item in original_structure.paragraphs
                                  if item.kind not in {"FIGURE_CAPTION", "TABLE_CAPTION"})
    protected_after = " ".join(item.text for item in final_structure.paragraphs
                                 if item.kind not in {"FIGURE_CAPTION", "TABLE_CAPTION"})
    all_text_before = " ".join(original_xml.xpath(".//w:t/text()", namespaces=NS))
    all_text_after = " ".join(final_xml.xpath(".//w:t/text()", namespaces=NS))
    checks = {"text_unchanged": texts_equal,
              "caption_notation_changes_authorized": not legacy.checks["unexpected_text_changes"],
              "numbers_unchanged": numeric_tokens(protected_before) == numeric_tokens(protected_after),
              "citations_unchanged": citation_tokens(all_text_before) == citation_tokens(all_text_after),
              "doi_unchanged": Counter(re.findall(r"\b10\.\d{4,9}/[^\s,;)>\]]+", all_text_before, re.I)) == Counter(re.findall(r"\b10\.\d{4,9}/[^\s,;)>\]]+", all_text_after, re.I)),
              "equations_unchanged": equation_xml_equal, "figures_unchanged": figure_xml_equal,
              "tables_unchanged": legacy.checks["tables_unchanged"], "references_unchanged": legacy.checks["reference_items_unchanged"],
              "headers_footers_and_other_parts_unchanged": unchanged_parts,
              "section_break_count_unchanged": legacy.checks["section_break_count_unchanged"]}
    return {"passed": all(checks.values()) and legacy.passed, "checks": checks,
            "warnings": ["Inspect Word pagination and visual layout before distribution."]}


def compliance_report_xlsx(*, journal: str, manuscript_name: str, template_version: int,
                           findings: list[Finding], approved_ids: set[str], integrity: dict[str, object] | None = None,
                           review_decisions: dict[str, dict[str, str]] | None = None) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Compliance"
    sheet.append(["Template compliance report"])
    sheet.append(["Journal", journal])
    sheet.append(["Manuscript", manuscript_name])
    sheet.append(["Article template version", template_version])
    sheet.append(["Compliance score", compliance_score(findings) / 100])
    sheet["B5"].number_format = "0%"
    sheet.append(["Scientific quality is not evaluated by this score."])
    sheet.append([])
    sheet.append(["Category", "Check", "Expected", "Detected", "Status", "Action applied", "Location",
                  "Protected-object type", "Preservation reason"])
    decisions = review_decisions or {}
    for finding in findings:
        decision = decisions.get(finding.id, {})
        action = (decision.get("status", "").replace("RESOLVED_", "").replace("_", " ").title()
                  if decision else
                  "Preserved without modification" if finding.status == "COMPLIANT_PROTECTED" else
                  "Preserved; warning reported" if finding.status == "WARNING_PRESERVED" else
                  "Fixed" if finding.id in approved_ids else
                  "Manual review" if finding.status not in {"COMPLIANT", "COMPLIANT_PROTECTED", "SAFE_FIX_AVAILABLE"} else "Not applied")
        sheet.append([finding.category, finding.check, finding.expected, finding.detected, finding.status,
                      action,
                      finding.target or "Document", finding.protected_object_type or "",
                      finding.preservation_reason or ""])
    if integrity:
        sheet.append([])
        sheet.append(["Integrity check", "PASS" if integrity["passed"] else "FAIL"])
        for key, result in integrity["checks"].items():
            sheet.append([key.replace("_", " ").title(), "PASS" if result else "FAIL"])
    for cell in sheet[8]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True)
    for col, width in {"A": 27, "B": 29, "C": 28, "D": 28, "E": 27, "F": 27, "G": 25,
                       "H": 28, "I": 55}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A9"
    sheet.auto_filter.ref = f"A8:I{8 + len(findings)}"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
