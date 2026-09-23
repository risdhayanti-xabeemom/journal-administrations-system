"""Formatting-preserving, paragraph-targeted Word patches and integrity checks."""

from __future__ import annotations

import copy
import io
import json
import re
from collections import Counter
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree

from services.revision_analysis import NS, W, analyze_manuscript, citation_tokens, numeric_tokens


class UnsafeRevisionError(ValueError):
    pass


HIGHLIGHT_COLORS = {"yellow": "yellow", "green": "green", "cyan": "cyan", "magenta": "magenta",
                    "red": "red", "blue": "blue", "darkyellow": "darkYellow", "darkgreen": "darkGreen"}


@dataclass(frozen=True)
class ApprovedPatch:
    paragraph_id: str
    original_text: str
    proposed_text: str
    numeric_approval_reason: str | None = None


@dataclass(frozen=True)
class IntegrityReport:
    passed: bool
    checks: dict[str, object]
    warnings: tuple[str, ...]

    def as_json(self) -> str:
        return json.dumps(asdict(self), ensure_ascii=False, indent=2)


def _xml(data: bytes) -> etree._Element:
    return etree.fromstring(data, parser=etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False))


def _paragraph_nodes(paragraph: etree._Element) -> list[etree._Element]:
    nodes = paragraph.xpath(".//w:t", namespaces=NS)
    for node in nodes:
        run = node.getparent()
        if run is None or run.tag != f"{{{W}}}r" or run.getparent() is not paragraph:
            raise UnsafeRevisionError("Target paragraph contains nested or tracked text that cannot be safely patched.")
        if len(run.xpath("./w:t", namespaces=NS)) != 1 or any(child.tag not in {f"{{{W}}}rPr", f"{{{W}}}t"} for child in run):
            raise UnsafeRevisionError("Target paragraph contains complex runs that require manual Word editing.")
    if not nodes:
        raise UnsafeRevisionError("Target paragraph has no patchable text runs.")
    return nodes


def _set_text(node: etree._Element, value: str) -> None:
    node.text = value
    if value.startswith(" ") or value.endswith(" "):
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _new_run(source_run: etree._Element, value: str, highlight_color: str | None) -> etree._Element:
    run = etree.Element(f"{{{W}}}r")
    source_properties = source_run.find(f"{{{W}}}rPr")
    properties = copy.deepcopy(source_properties) if source_properties is not None else etree.Element(f"{{{W}}}rPr")
    if highlight_color:
        old_highlight = properties.find(f"{{{W}}}highlight")
        if old_highlight is not None:
            properties.remove(old_highlight)
        marker = etree.SubElement(properties, f"{{{W}}}highlight")
        marker.set(f"{{{W}}}val", highlight_color)
    if len(properties) or source_properties is not None:
        run.append(properties)
    node = etree.SubElement(run, f"{{{W}}}t")
    _set_text(node, value)
    return run


def _replace_span(paragraph: etree._Element, start: int, end: int, replacement: str, highlight_color: str | None) -> None:
    nodes = _paragraph_nodes(paragraph)
    offsets: list[tuple[etree._Element, int, int]] = []
    cursor = 0
    for node in nodes:
        length = len(node.text or "")
        offsets.append((node, cursor, cursor + length))
        cursor += length
    if start < 0 or end < start or end > cursor:
        raise UnsafeRevisionError("Patch text span is outside the original paragraph.")
    if start == end:
        selected = next(((node, lo, hi) for node, lo, hi in offsets if lo <= start <= hi), offsets[-1])
        first = last = selected
    else:
        overlaps = [entry for entry in offsets if entry[2] > start and entry[1] < end]
        if not overlaps:
            raise UnsafeRevisionError("Patch span cannot be mapped to Word text runs.")
        first, last = overlaps[0], overlaps[-1]
    first_node, first_start, _ = first
    last_node, last_start, _ = last
    first_run = first_node.getparent()
    original_first = first_node.text or ""
    prefix = original_first[:start - first_start]
    suffix = (last_node.text or "")[end - last_start:]
    _set_text(first_node, prefix)
    if first_node is last_node:
        if replacement:
            first_run.addnext(_new_run(first_run, replacement, highlight_color))
        if suffix:
            anchor = first_run.getnext() if replacement else first_run
            anchor.addnext(_new_run(first_run, suffix, None))
    else:
        for node, _, _ in offsets[offsets.index(first) + 1:offsets.index(last)]:
            _set_text(node, "")
        _set_text(last_node, suffix)
        if replacement:
            first_run.addnext(_new_run(first_run, replacement, highlight_color))


def patch_manuscript(content: bytes, patches: list[ApprovedPatch], *, highlight: bool, highlight_color: str = "yellow") -> bytes:
    if highlight and highlight_color.lower() not in HIGHLIGHT_COLORS:
        raise UnsafeRevisionError("Unsupported revision highlight color. Use a Word-compatible configured color.")
    structure = analyze_manuscript(content)
    by_id = {paragraph.identifier: paragraph for paragraph in structure.paragraphs}
    if len({patch.paragraph_id for patch in patches}) != len(patches):
        raise UnsafeRevisionError("More than one approved proposal targets the same paragraph. Merge them explicitly first.")
    with ZipFile(io.BytesIO(content)) as source:
        root = _xml(source.read("word/document.xml"))
        paragraphs = root.xpath(".//w:body//w:p", namespaces=NS)
        for patch in patches:
            paragraph = by_id.get(patch.paragraph_id)
            if not paragraph or not paragraph.patchable or paragraph.text != patch.original_text:
                raise UnsafeRevisionError(f"{patch.paragraph_id} is missing, changed, or contains protected content.")
            if citation_tokens(patch.original_text) != citation_tokens(patch.proposed_text):
                raise UnsafeRevisionError(f"Citation identifiers changed in {patch.paragraph_id}.")
            if numeric_tokens(patch.original_text) != numeric_tokens(patch.proposed_text) and not patch.numeric_approval_reason:
                raise UnsafeRevisionError(f"Unauthorized numerical change in {patch.paragraph_id}.")
            target = paragraphs[paragraph.index - 1]
            _paragraph_nodes(target)
            operations = SequenceMatcher(None, patch.original_text, patch.proposed_text, autojunk=False).get_opcodes()
            for operation, start, end, new_start, new_end in reversed(operations):
                if operation != "equal":
                    _replace_span(target, start, end, patch.proposed_text[new_start:new_end],
                                  HIGHLIGHT_COLORS[highlight_color.lower()] if highlight else None)
        output = io.BytesIO()
        with ZipFile(output, "w", ZIP_DEFLATED) as destination:
            for item in source.infolist():
                data = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True) if item.filename == "word/document.xml" else source.read(item.filename)
                destination.writestr(item, data)
    return output.getvalue()


def _package_inventory(content: bytes) -> tuple[dict[str, bytes], etree._Element]:
    with ZipFile(io.BytesIO(content)) as archive:
        protected = {
            name: archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/") or re.fullmatch(r"word/(header|footer)\d+\.xml", name)
        }
        return protected, _xml(archive.read("word/document.xml"))


def check_integrity(original: bytes, revised: bytes, patches: list[ApprovedPatch]) -> IntegrityReport:
    original_structure = analyze_manuscript(original)
    revised_structure = analyze_manuscript(revised)
    before_parts, before_xml = _package_inventory(original)
    after_parts, after_xml = _package_inventory(revised)
    original_paragraphs = {para.identifier: para for para in original_structure.paragraphs}
    revised_paragraphs = {para.identifier: para for para in revised_structure.paragraphs}
    patch_map = {patch.paragraph_id: patch for patch in patches}
    unexpected_text = []
    unauthorized_numbers = []
    approved_numeric_changes = []
    for identifier, before in original_paragraphs.items():
        after = revised_paragraphs.get(identifier)
        if after is None:
            unexpected_text.append(f"Missing {identifier}")
            continue
        expected = patch_map[identifier].proposed_text if identifier in patch_map else before.text
        if after.text != expected:
            unexpected_text.append(identifier)
        if numeric_tokens(before.text) != numeric_tokens(after.text):
            difference = {"paragraph": identifier, "original": before.text, "revised": after.text}
            if identifier in patch_map and patch_map[identifier].numeric_approval_reason:
                approved_numeric_changes.append({**difference, "approval_reason": patch_map[identifier].numeric_approval_reason})
            else:
                unauthorized_numbers.append(difference)
    original_citations = citation_tokens("\n".join(para.text for para in original_structure.paragraphs))
    revised_citations = citation_tokens("\n".join(para.text for para in revised_structure.paragraphs))
    section_original = [para.text for para in original_structure.paragraphs if para.kind == "SECTION_HEADING"]
    section_revised = [para.text for para in revised_structure.paragraphs if para.kind == "SECTION_HEADING"]
    checks: dict[str, object] = {
        "paragraph_count_unchanged": len(original_structure.paragraphs) == len(revised_structure.paragraphs),
        "section_headings_unchanged": section_original == section_revised,
        "section_break_count_unchanged": len(before_xml.xpath(".//w:sectPr", namespaces=NS)) == len(after_xml.xpath(".//w:sectPr", namespaces=NS)),
        "protected_package_parts_unchanged": before_parts == after_parts,
        "tables_unchanged": original_structure.table_count == revised_structure.table_count and
            ["".join(table.itertext()) for table in before_xml.xpath(".//w:tbl", namespaces=NS)] ==
            ["".join(table.itertext()) for table in after_xml.xpath(".//w:tbl", namespaces=NS)],
        "figures_unchanged": original_structure.figure_count == revised_structure.figure_count,
        "equations_unchanged": original_structure.equation_count == revised_structure.equation_count,
        "table_captions_present": len(original_structure.table_captions) == len(revised_structure.table_captions),
        "figure_captions_present": len(original_structure.figure_captions) == len(revised_structure.figure_captions),
        "reference_items_unchanged": original_structure.references == revised_structure.references,
        "citation_labels_unchanged": original_citations == revised_citations,
        "unexpected_text_changes": unexpected_text,
        "unauthorized_numeric_changes": unauthorized_numbers,
        "approved_numeric_changes": approved_numeric_changes,
    }
    passed = all(value is True for key, value in checks.items() if isinstance(value, bool)) and not unexpected_text and not unauthorized_numbers
    warnings = ("Page count and line wrapping require visual DOCX rendering before external distribution.",)
    return IntegrityReport(passed, checks, warnings)


def response_to_reviewers(article_title: str, submission_id: str, revision_round: int, comments: list[dict[str, str]], *, draft: bool = False) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    document.add_heading("Draft Response to Reviewers - Not Final" if draft else "Response to Reviewers", level=0)
    document.add_paragraph(f"Article title: {article_title}")
    document.add_paragraph(f"Submission ID: {submission_id or 'Standalone revision'}")
    document.add_paragraph(f"Revision round: {revision_round}")
    current_reviewer = None
    for item in comments:
        if item["reviewer"] != current_reviewer:
            current_reviewer = item["reviewer"]
            document.add_heading(current_reviewer, level=1)
        document.add_heading(f"Comment {item['number']}", level=2)
        for label, value in (
            ("Reviewer Comment", item["raw_text"]),
            ("Response", item["response"]),
            ("Revision Made", item["revision_made"]),
            ("Location", item["location"]),
            ("Status", item["status"]),
        ):
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
